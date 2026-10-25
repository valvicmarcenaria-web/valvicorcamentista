#!/usr/bin/env python3
"""Gera o modelo de planilha financeira da Valvic (xlsx com fórmulas vivas).

Abas: Premissas, Plano_Contas, Lancamentos, Projetos, Fluxo_Mensal, DRE,
Recebiveis, Checks. Fluxo e DRE saem da MESMA base de lançamentos:
  - Fluxo_Mensal soma por `data_caixa` (regime de caixa)
  - DRE soma por `data_competencia` (regime de competência)

Regra de lançamento para competência em marcenaria:
  - recebimentos de cliente (sinal, parcelas) e custos variáveis do projeto
    (chapa, ferragem, terceirizados, comissões, RT) -> data_competencia = data
    de ENTREGA do projeto; data_caixa = data do pagamento/vencimento.
  - fixos -> mês de uso/trabalho.

Uso:
  python gerar_planilha_financeira.py saida.xlsx [--inicio 2026-01]
         [--meses 12] [--saldo-inicial 0] [--lancamentos base.csv] [--exemplo]

  --lancamentos: CSV/XLSX com colunas data_competencia, data_caixa, status,
                 descricao, categoria, valor, projeto, contraparte, conta,
                 documento (as que faltarem ficam vazias).
  --exemplo:     preenche lançamentos e projetos fictícios para demonstração.
"""
import argparse
import datetime as dt
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

MAX_LINHAS = 5000  # alcance das fórmulas sobre Lancamentos

AZUL = Font(color="1F4E9E")
NEGRITO = Font(bold=True)
TITULO = Font(bold=True, size=14)
CAB = PatternFill("solid", fgColor="DDE4EE")
SUB = PatternFill("solid", fgColor="F2F2F2")
VERMELHO = PatternFill("solid", fgColor="F8D7DA")
VERDE = PatternFill("solid", fgColor="D4EDDA")
MOEDA = 'R$ #,##0.00;[Red]-R$ #,##0.00'
PCT = '0.0%'
DATA = 'DD/MM/YYYY'
MES = 'MMM/YY'

# (categoria, grupo_dre, grupo_fluxo, natureza)
PLANO = [
    ("Entrada/sinal de cliente", "Receita bruta", "Recebimentos de clientes", "receita"),
    ("Parcela de cliente", "Receita bruta", "Recebimentos de clientes", "receita"),
    ("Recebimento cartão (líquido)", "Receita bruta", "Recebimentos de clientes", "receita"),
    ("Outras receitas", "Receita bruta", "Outras entradas", "receita"),
    ("Impostos sobre NF (DAS)", "Deduções", "Impostos sobre faturamento", "variavel"),
    ("Chapas e fita", "Custos variáveis do projeto", "Material", "variavel"),
    ("Ferragens", "Custos variáveis do projeto", "Material", "variavel"),
    ("Iluminação / LED", "Custos variáveis do projeto", "Material", "variavel"),
    ("Vidros e espelhos", "Custos variáveis do projeto", "Material", "variavel"),
    ("Insumos (cola, parafusos, embalagem)", "Custos variáveis do projeto", "Material", "variavel"),
    ("Terceirizados", "Custos variáveis do projeto", "Terceirizados", "variavel"),
    ("Frete / carreto", "Custos variáveis do projeto", "Logística", "variavel"),
    ("Comissão vendedor", "Custos variáveis de venda", "Comissões e RT", "variavel"),
    ("Comissão produção / montagem", "Custos variáveis de venda", "Comissões e RT", "variavel"),
    ("RT arquiteto", "Custos variáveis de venda", "Comissões e RT", "variavel"),
    ("Taxa maquininha", "Custos variáveis de venda", "Comissões e RT", "variavel"),
    ("Salários e encargos", "Custos fixos de produção", "Folha e pró-labore", "fixo"),
    ("Pró-labore", "Despesas fixas adm/comerciais", "Folha e pró-labore", "fixo"),
    ("Aluguel galpão", "Despesas fixas adm/comerciais", "Despesas fixas", "fixo"),
    ("Energia / água / internet", "Custos fixos de produção", "Despesas fixas", "fixo"),
    ("Manutenção de máquinas", "Custos fixos de produção", "Despesas fixas", "fixo"),
    ("Sistemas e contador", "Despesas fixas adm/comerciais", "Despesas fixas", "fixo"),
    ("Marketing", "Despesas fixas adm/comerciais", "Despesas fixas", "fixo"),
    ("Outras despesas fixas", "Despesas fixas adm/comerciais", "Despesas fixas", "fixo"),
    ("Juros e tarifas", "Resultado financeiro", "Despesas fixas", "financeiro"),
    ("Parcela financiamento máquinas", "(fora da DRE)", "Dívidas e sócios", "nao_operacional"),
    ("Empréstimo sócio (entrada)", "(fora da DRE)", "Dívidas e sócios", "nao_operacional"),
    ("Devolução empréstimo sócio", "(fora da DRE)", "Dívidas e sócios", "nao_operacional"),
    ("Investimento em ativo", "(fora da DRE)", "Investimentos", "nao_operacional"),
    ("Transferência entre contas", "(fora da DRE)", "Transferências", "nao_operacional"),
]

GRUPOS_FLUXO = [
    ("ENTRADAS", ["Recebimentos de clientes", "Outras entradas"]),
    ("SAÍDAS VARIÁVEIS", ["Material", "Terceirizados", "Logística", "Comissões e RT",
                          "Impostos sobre faturamento"]),
    ("SAÍDAS FIXAS", ["Folha e pró-labore", "Despesas fixas"]),
    ("INVESTIMENTOS E FINANCIAMENTOS", ["Investimentos", "Dívidas e sócios", "Transferências"]),
]

PREMISSAS = [
    ("Saldo inicial de caixa (todas as contas)", 0, MOEDA),
    ("Saldo mínimo / reserva-alvo", 70000, MOEDA),
    ("NF sobre faturamento", 0.07, PCT),
    ("Taxa maquininha / parcelamento", 0.07, PCT),
    ("Comissão vendedor", 0.05, PCT),
    ("Comissão produção", 0.05, PCT),
    ("RT arquiteto (sobre líquido)", 0.10, PCT),
    ("Entrada mínima no contrato", 0.40, PCT),
    ("MC% piso (break-even)", 0.43, PCT),
    ("MC% alvo saudável", 0.48, PCT),
    ("Custo fixo mensal de referência", 47219, MOEDA),
    ("Parcelas de dívidas por mês (a levantar)", 0, MOEDA),
    ("Materialidade de auditoria por lançamento", 200, MOEDA),
]

LANC_COLS = ["data_competencia", "data_caixa", "status", "descricao", "categoria",
             "grupo_dre", "grupo_fluxo", "valor", "projeto", "contraparte", "conta",
             "documento"]


def col(name):
    return get_column_letter(LANC_COLS.index(name) + 1)


def rng(name):
    c = col(name)
    return f"Lancamentos!${c}$2:${c}${MAX_LINHAS}"


def add_months(d, n):
    m = d.month - 1 + n
    return dt.date(d.year + m // 12, m % 12 + 1, 1)


def largura(ws, larguras):
    for i, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def cabecalho(ws, linha, valores):
    for i, v in enumerate(valores, 1):
        c = ws.cell(linha, i, v)
        c.font = NEGRITO
        c.fill = CAB
        c.alignment = Alignment(wrap_text=True, vertical="center")


def aba_premissas(wb):
    ws = wb.active
    ws.title = "Premissas"
    ws["A1"] = "Premissas — editar só as células em azul"
    ws["A1"].font = TITULO
    cabecalho(ws, 3, ["Premissa", "Valor"])
    refs = {}
    for i, (nome, valor, fmt) in enumerate(PREMISSAS, 4):
        ws.cell(i, 1, nome)
        c = ws.cell(i, 2, valor)
        c.font = AZUL
        c.number_format = fmt
        refs[nome] = f"Premissas!$B${i}"
    largura(ws, [46, 18])
    return refs


def aba_plano(wb):
    ws = wb.create_sheet("Plano_Contas")
    cabecalho(ws, 1, ["categoria", "grupo_dre", "grupo_fluxo", "natureza"])
    for i, linha in enumerate(PLANO, 2):
        for j, v in enumerate(linha, 1):
            ws.cell(i, j, v)
    largura(ws, [38, 32, 30, 16])
    ws.freeze_panes = "A2"
    return f"Plano_Contas!$A$2:$A${len(PLANO) + 1}"


def aba_lancamentos(wb, linhas, lista_categorias):
    ws = wb.create_sheet("Lancamentos")
    cabecalho(ws, 1, LANC_COLS)
    ultima = len(PLANO) + 1
    for r in range(2, max(len(linhas) + 2, 3)):
        dados = linhas[r - 2] if r - 2 < len(linhas) else {}
        for nome in LANC_COLS:
            c = ws[f"{col(nome)}{r}"]
            if nome == "grupo_dre":
                c.value = (f'=IF({col("categoria")}{r}="","",IFERROR(VLOOKUP({col("categoria")}{r},'
                           f'Plano_Contas!$A$2:$C${ultima},2,FALSE),"SEM CATEGORIA"))')
            elif nome == "grupo_fluxo":
                c.value = (f'=IF({col("categoria")}{r}="","",IFERROR(VLOOKUP({col("categoria")}{r},'
                           f'Plano_Contas!$A$2:$C${ultima},3,FALSE),"SEM CATEGORIA"))')
            else:
                v = dados.get(nome)
                c.value = v if v not in ("", None) else None
                c.font = AZUL
            if nome.startswith("data"):
                c.number_format = DATA
            if nome == "valor":
                c.number_format = MOEDA
    dv_cat = DataValidation(type="list", formula1=f"={lista_categorias}", allow_blank=True)
    dv_st = DataValidation(type="list", formula1='"pago,aberto,previsto"', allow_blank=True)
    ws.add_data_validation(dv_cat)
    ws.add_data_validation(dv_st)
    dv_cat.add(f"{col('categoria')}2:{col('categoria')}{MAX_LINHAS}")
    dv_st.add(f"{col('status')}2:{col('status')}{MAX_LINHAS}")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(LANC_COLS))}{max(len(linhas) + 1, 2)}"
    largura(ws, [13, 13, 9, 34, 32, 26, 24, 14, 22, 22, 12, 14])


def aba_projetos(wb, projetos):
    ws = wb.create_sheet("Projetos")
    cols = ["projeto", "cliente", "valor_contrato", "data_assinatura", "data_entrega",
            "mc_orcada_pct", "recebido", "a_receber", "custo_variavel_realizado",
            "mc_realizada_R$", "mc_realizada_pct", "desvio_mc_pp"]
    cabecalho(ws, 1, cols)
    n = max(len(projetos), 15)
    proj, val, st, cat, dre = (rng("projeto"), rng("valor"), rng("status"),
                               rng("categoria"), rng("grupo_dre"))
    for r in range(2, n + 2):
        p = projetos[r - 2] if r - 2 < len(projetos) else {}
        for j, k in enumerate(cols[:6], 1):
            c = ws.cell(r, j, p.get(k))
            c.font = AZUL
        ws[f"C{r}"].number_format = MOEDA
        ws[f"D{r}"].number_format = DATA
        ws[f"E{r}"].number_format = DATA
        ws[f"F{r}"].number_format = PCT
        ws[f"G{r}"] = f'=IF(A{r}="","",SUMIFS({val},{proj},A{r},{dre},"Receita bruta",{st},"pago"))'
        ws[f"H{r}"] = f'=IF(A{r}="","",C{r}-G{r})'
        ws[f"I{r}"] = (f'=IF(A{r}="","",-SUMIFS({val},{proj},A{r},{dre},"Custos variáveis do projeto")'
                       f'-SUMIFS({val},{proj},A{r},{dre},"Custos variáveis de venda")'
                       f'-SUMIFS({val},{proj},A{r},{dre},"Deduções"))')
        ws[f"J{r}"] = f'=IF(A{r}="","",C{r}-I{r})'
        ws[f"K{r}"] = f'=IF(OR(A{r}="",C{r}=0),"",J{r}/C{r})'
        ws[f"L{r}"] = f'=IF(OR(K{r}="",F{r}=""),"",(K{r}-F{r})*100)'
        for c in "GHIJ":
            ws[f"{c}{r}"].number_format = MOEDA
        ws[f"K{r}"].number_format = PCT
        ws[f"L{r}"].number_format = '0.0'
    ws.conditional_formatting.add(f"L2:L{n + 1}", CellIsRule(operator="lessThan", formula=["-3"], fill=VERMELHO))
    ws["N1"] = "MC realizada só vale para projeto ENTREGUE com todos os custos lançados."
    ws["N2"] = "Desvio < -3 p.p. = investigar (preço de chapa, retrabalho, item esquecido, desconto)."
    ws.freeze_panes = "B2"
    largura(ws, [22, 20, 15, 13, 13, 11, 14, 14, 16, 15, 11, 10])
    return n


def aba_fluxo(wb, inicio, meses, prem):
    ws = wb.create_sheet("Fluxo_Mensal")
    ws["A1"] = "Fluxo de caixa mensal — regime de CAIXA (soma por data_caixa; inclui pago + aberto + previsto)"
    ws["A1"].font = TITULO
    ws["A3"] = "Mês"
    ws["A3"].font = NEGRITO
    for m in range(meses):
        c = ws.cell(3, 2 + m, add_months(inicio, m))
        c.number_format = MES
        c.font = NEGRITO
        c.fill = CAB
    ultima_col = get_column_letter(1 + meses)
    ws.cell(3, 2 + meses, "Total").font = NEGRITO
    tot_col = get_column_letter(2 + meses)

    val, gf, dc = rng("valor"), rng("grupo_fluxo"), rng("data_caixa")
    r = 4
    ws.cell(r, 1, "SALDO INICIAL").font = NEGRITO
    linha_saldo_ini = r
    r += 2
    linhas_total_bloco = []
    for bloco, grupos in GRUPOS_FLUXO:
        ws.cell(r, 1, bloco).font = NEGRITO
        ws.cell(r, 1).fill = SUB
        r += 1
        ini = r
        for g in grupos:
            ws.cell(r, 1, "   " + g)
            for m in range(meses):
                L = get_column_letter(2 + m)
                ws[f"{L}{r}"] = (f'=SUMIFS({val},{gf},"{g}",{dc},">="&{L}$3,'
                                 f'{dc},"<="&EOMONTH({L}$3,0))')
                ws[f"{L}{r}"].number_format = MOEDA
            ws[f"{tot_col}{r}"] = f"=SUM(B{r}:{ultima_col}{r})"
            ws[f"{tot_col}{r}"].number_format = MOEDA
            r += 1
        ws.cell(r, 1, f"Total {bloco.lower()}").font = NEGRITO
        for m in range(meses + 1):
            L = get_column_letter(2 + m)
            ws[f"{L}{r}"] = f"=SUM({L}{ini}:{L}{r - 1})"
            ws[f"{L}{r}"].number_format = MOEDA
            ws[f"{L}{r}"].font = NEGRITO
        linhas_total_bloco.append(r)
        r += 2
    ws.cell(r, 1, "SEM CATEGORIA (corrigir!)")
    for m in range(meses):
        L = get_column_letter(2 + m)
        ws[f"{L}{r}"] = f'=SUMIFS({val},{gf},"SEM CATEGORIA",{dc},">="&{L}$3,{dc},"<="&EOMONTH({L}$3,0))'
        ws[f"{L}{r}"].number_format = MOEDA
    linha_semcat = r
    r += 2
    ws.cell(r, 1, "VARIAÇÃO DO CAIXA").font = NEGRITO
    linha_var = r
    soma_blocos = "+".join(f"{{L}}{x}" for x in linhas_total_bloco) + f"+{{L}}{linha_semcat}"
    for m in range(meses):
        L = get_column_letter(2 + m)
        ws[f"{L}{r}"] = "=" + soma_blocos.format(L=L)
        ws[f"{L}{r}"].number_format = MOEDA
    r += 1
    ws.cell(r, 1, "SALDO FINAL").font = NEGRITO
    linha_saldo = r
    for m in range(meses):
        L = get_column_letter(2 + m)
        ws[f"{L}{linha_saldo_ini}"] = (f"={prem['Saldo inicial de caixa (todas as contas)']}" if m == 0
                                      else f"={get_column_letter(1 + m)}{linha_saldo}")
        ws[f"{L}{linha_saldo_ini}"].number_format = MOEDA
        ws[f"{L}{r}"] = f"={L}{linha_saldo_ini}+{L}{linha_var}"
        ws[f"{L}{r}"].number_format = MOEDA
        ws[f"{L}{r}"].font = NEGRITO
    r += 1
    ws.cell(r, 1, "Abaixo da reserva-alvo?")
    for m in range(meses):
        L = get_column_letter(2 + m)
        ws[f"{L}{r}"] = f'=IF({L}{linha_saldo}<0,"NEGATIVO",IF({L}{linha_saldo}<{prem["Saldo mínimo / reserva-alvo"]},"abaixo","ok"))'
    faixa = f"B{linha_saldo}:{ultima_col}{linha_saldo}"
    ws.conditional_formatting.add(faixa, CellIsRule(operator="lessThan", formula=["0"], fill=VERMELHO))
    ws.conditional_formatting.add(faixa, CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=VERDE))
    ws.freeze_panes = "B4"
    largura(ws, [38] + [14] * (meses + 1))
    return linha_semcat


DRE_LINHAS = [
    ("Receita bruta", "Receita bruta", 1),
    ("(−) Deduções (impostos s/ NF)", "Deduções", 1),
    ("(=) RECEITA LÍQUIDA", "=RL", 0),
    ("(−) Custos variáveis do projeto", "Custos variáveis do projeto", 1),
    ("(−) Custos variáveis de venda", "Custos variáveis de venda", 1),
    ("(=) MARGEM DE CONTRIBUIÇÃO", "=MC", 0),
    ("   MC % da receita bruta", "=MCP", 0),
    ("(−) Custos fixos de produção", "Custos fixos de produção", 1),
    ("(−) Despesas fixas adm/comerciais", "Despesas fixas adm/comerciais", 1),
    ("(=) RESULTADO OPERACIONAL", "=RO", 0),
    ("(−/+) Resultado financeiro", "Resultado financeiro", 1),
    ("(=) RESULTADO LÍQUIDO (antes de depreciação)", "=RLQ", 0),
    ("Fora da DRE (dívidas, sócios, investimentos) — só caixa", "(fora da DRE)", 2),
]


def aba_dre(wb, inicio, meses):
    ws = wb.create_sheet("DRE")
    ws["A1"] = "DRE gerencial — regime de COMPETÊNCIA (soma por data_competencia)"
    ws["A1"].font = TITULO
    ws["A2"] = "Receita e custos variáveis do projeto entram no mês da ENTREGA (data_competencia)."
    for m in range(meses):
        c = ws.cell(3, 2 + m, add_months(inicio, m))
        c.number_format = MES
        c.font = NEGRITO
        c.fill = CAB
    ws.cell(3, 2 + meses, "Total").font = NEGRITO
    val, gd, dcomp = rng("valor"), rng("grupo_dre"), rng("data_competencia")
    pos = {}
    for i, (rot, chave, tipo) in enumerate(DRE_LINHAS, 4):
        ws.cell(i, 1, rot)
        if tipo != 1:
            ws.cell(i, 1).font = NEGRITO
        pos[chave] = i
        for m in range(meses + 1):
            L = get_column_letter(2 + m)
            cel = ws[f"{L}{i}"]
            if m == meses:
                cel.value = (f"=IFERROR({L}{pos['=MC']}/{L}{pos['Receita bruta']},0)" if chave == "=MCP"
                             else f"=SUM(B{i}:{get_column_letter(1 + meses)}{i})")
            elif tipo in (1, 2):
                cel.value = f'=SUMIFS({val},{gd},"{chave}",{dcomp},">="&{L}$3,{dcomp},"<="&EOMONTH({L}$3,0))'
            elif chave == "=RL":
                cel.value = f"={L}{pos['Receita bruta']}+{L}{pos['Deduções']}"
            elif chave == "=MC":
                cel.value = (f"={L}{pos['=RL']}+{L}{pos['Custos variáveis do projeto']}"
                             f"+{L}{pos['Custos variáveis de venda']}")
            elif chave == "=MCP":
                cel.value = f"=IFERROR({L}{pos['=MC']}/{L}{pos['Receita bruta']},0)"
            elif chave == "=RO":
                cel.value = (f"={L}{pos['=MC']}+{L}{pos['Custos fixos de produção']}"
                             f"+{L}{pos['Despesas fixas adm/comerciais']}")
            elif chave == "=RLQ":
                cel.value = f"={L}{pos['=RO']}+{L}{pos['Resultado financeiro']}"
            cel.number_format = PCT if chave == "=MCP" else MOEDA
            if tipo == 0:
                cel.font = NEGRITO
    ws.freeze_panes = "B4"
    largura(ws, [52] + [14] * (meses + 1))


def aba_recebiveis(wb, n_proj):
    ws = wb.create_sheet("Recebiveis")
    ws["A1"] = "Aging de recebíveis — lançamentos de receita com status aberto/previsto (base: data_caixa = vencimento)"
    ws["A1"].font = NEGRITO
    cols = ["cliente/projeto", "a vencer", "1–30 d", "31–60 d", "61–90 d", "> 90 d", "total aberto"]
    cabecalho(ws, 3, cols)
    val, proj, st, gd, dc = rng("valor"), rng("projeto"), rng("status"), rng("grupo_dre"), rng("data_caixa")
    base = f'{val},{proj},$A{{r}},{gd},"Receita bruta",{st},"<>pago"'
    faixas = [
        f'{dc},">="&TODAY()',
        f'{dc},"<"&TODAY(),{dc},">="&(TODAY()-30)',
        f'{dc},"<"&(TODAY()-30),{dc},">="&(TODAY()-60)',
        f'{dc},"<"&(TODAY()-60),{dc},">="&(TODAY()-90)',
        f'{dc},"<"&(TODAY()-90)',
    ]
    for i in range(n_proj):
        r = 4 + i
        ws[f"A{r}"] = f'=IF(Projetos!A{2 + i}="","",Projetos!A{2 + i})'
        for j, f in enumerate(faixas):
            L = get_column_letter(2 + j)
            ws[f"{L}{r}"] = f'=IF($A{r}="","",SUMIFS({base.format(r=r)},{f}))'
            ws[f"{L}{r}"].number_format = MOEDA
        ws[f"G{r}"] = f'=IF($A{r}="","",SUM(B{r}:F{r}))'
        ws[f"G{r}"].number_format = MOEDA
    r = 4 + n_proj
    ws[f"A{r}"] = "TOTAL"
    ws[f"A{r}"].font = NEGRITO
    for j in range(6):
        L = get_column_letter(2 + j)
        ws[f"{L}{r}"] = f"=SUM({L}4:{L}{r - 1})"
        ws[f"{L}{r}"].number_format = MOEDA
        ws[f"{L}{r}"].font = NEGRITO
    ws.conditional_formatting.add(f"D4:F{r - 1}", CellIsRule(operator="greaterThan", formula=["0"], fill=VERMELHO))
    largura(ws, [26, 14, 14, 14, 14, 14, 15])


def aba_checks(wb, linha_semcat, meses):
    ws = wb.create_sheet("Checks")
    ws["A1"] = "Checks de integridade — todos devem dar 0 (ou OK)"
    ws["A1"].font = TITULO
    cabecalho(ws, 3, ["Teste", "Resultado", "Leitura"])
    cat, val, dcomp, dc, st = (rng("categoria"), rng("valor"), rng("data_competencia"),
                               rng("data_caixa"), rng("status"))
    gd = rng("grupo_dre")
    ult = get_column_letter(1 + meses)
    testes = [
        ("Lançamentos com categoria fora do plano", f'=COUNTIF({gd},"SEM CATEGORIA")',
         "corrigir a categoria no Lancamentos"),
        ("Lançamentos com valor e sem categoria", f'=COUNTIFS({val},"<>",{cat},"")', "classificar"),
        ("Lançamentos sem data_caixa", f'=COUNTIFS({val},"<>",{dc},"")', "sem data não entra no fluxo"),
        ("Lançamentos sem data_competencia", f'=COUNTIFS({val},"<>",{dcomp},"")', "sem data não entra na DRE"),
        ("Lançamentos sem status", f'=COUNTIFS({val},"<>",{st},"")', "pago / aberto / previsto"),
        ("Valor SEM CATEGORIA no fluxo", f"=SUM(Fluxo_Mensal!B{linha_semcat}:{ult}{linha_semcat})",
         "deve ser 0"),
        ("Σ lançamentos no período (caixa) − variação do fluxo",
         f'=ROUND(SUMIFS({val},{dc},">="&Fluxo_Mensal!B3,{dc},"<="&EOMONTH(Fluxo_Mensal!{ult}3,0))'
         f'-SUM(Fluxo_Mensal!B{linha_semcat + 2}:{ult}{linha_semcat + 2}),2)',
         "≠ 0 = fórmula do fluxo não pega algum grupo"),
        ("Receitas com valor negativo", f'=COUNTIFS({gd},"Receita bruta",{val},"<0")',
         "estorno? confirmar"),
        ("Custos/despesas com valor positivo",
         f'=COUNTIFS({gd},"Custos*",{val},">0")+COUNTIFS({gd},"Despesas*",{val},">0")',
         "saída deve ser negativa"),
    ]
    for i, (nome, formula, leitura) in enumerate(testes, 4):
        ws.cell(i, 1, nome)
        ws.cell(i, 2, formula)
        ws.cell(i, 3, leitura)
    ws.conditional_formatting.add(f"B4:B{3 + len(testes)}",
                                  FormulaRule(formula=["B4<>0"], fill=VERMELHO))
    ws.conditional_formatting.add(f"B4:B{3 + len(testes)}",
                                  FormulaRule(formula=["B4=0"], fill=VERDE))
    largura(ws, [52, 14, 44])


def ler_lancamentos(caminho):
    import pandas as pd
    if caminho.lower().endswith((".xlsx", ".xls")):
        df = pd.read_excel(caminho)
    else:
        try:
            df = pd.read_csv(caminho, sep=None, engine="python", encoding="utf-8-sig")
        except UnicodeDecodeError:
            df = pd.read_csv(caminho, sep=None, engine="python", encoding="latin-1")
    df.columns = [str(c).strip().lower() for c in df.columns]
    linhas = []
    for _, row in df.iterrows():
        d = {}
        for k in LANC_COLS:
            if k in df.columns and pd.notna(row[k]):
                v = row[k]
                if k.startswith("data"):
                    v = pd.to_datetime(v, dayfirst=True).date()
                elif k == "valor" and isinstance(v, str):
                    v = float(v.replace("R$", "").replace(".", "").replace(",", ".").strip())
                d[k] = v
        linhas.append(d)
    return linhas


def exemplo(inicio):
    d = lambda m, dia: dt.date(add_months(inicio, m).year, add_months(inicio, m).month, dia)  # noqa: E731
    projetos = [
        {"projeto": "Cozinha Ana", "cliente": "Ana", "valor_contrato": 40000,
         "data_assinatura": d(0, 5), "data_entrega": d(1, 25), "mc_orcada_pct": 0.45},
        {"projeto": "Closet Bruno", "cliente": "Bruno", "valor_contrato": 25000,
         "data_assinatura": d(0, 20), "data_entrega": d(2, 15), "mc_orcada_pct": 0.42},
    ]
    L = []
    def add(comp, caixa, st, desc, cat, valor, proj=""):
        L.append({"data_competencia": comp, "data_caixa": caixa, "status": st, "descricao": desc,
                  "categoria": cat, "valor": valor, "projeto": proj})
    ea, eb = d(1, 25), d(2, 15)
    add(ea, d(0, 5), "pago", "Entrada 40% Ana", "Entrada/sinal de cliente", 16000, "Cozinha Ana")
    add(ea, d(1, 25), "pago", "Parcela entrega Ana", "Parcela de cliente", 24000, "Cozinha Ana")
    add(ea, d(0, 18), "pago", "Chapas Ana", "Chapas e fita", -9800, "Cozinha Ana")
    add(ea, d(1, 10), "pago", "Ferragens Ana (28dd)", "Ferragens", -3600, "Cozinha Ana")
    add(ea, d(1, 20), "pago", "Vidraceiro Ana", "Terceirizados", -1500, "Cozinha Ana")
    add(ea, d(2, 20), "previsto", "DAS sobre NF Ana", "Impostos sobre NF (DAS)", -2800, "Cozinha Ana")
    add(ea, d(1, 28), "pago", "Comissão vendedor Ana", "Comissão vendedor", -2000, "Cozinha Ana")
    add(eb, d(0, 20), "pago", "Entrada 40% Bruno", "Entrada/sinal de cliente", 10000, "Closet Bruno")
    add(eb, d(2, 15), "aberto", "Parcela entrega Bruno", "Parcela de cliente", 15000, "Closet Bruno")
    add(eb, d(1, 5), "pago", "Chapas Bruno", "Chapas e fita", -6200, "Closet Bruno")
    add(eb, d(2, 5), "aberto", "Ferragens Bruno", "Ferragens", -2400, "Closet Bruno")
    for m in range(3):
        add(d(m, 28), d(m + 1, 5), "pago" if m < 2 else "previsto", "Folha", "Salários e encargos", -32000)
        add(d(m, 1), d(m, 10), "pago" if m < 2 else "previsto", "Aluguel galpão", "Aluguel galpão", -8000)
        add(d(m, 1), d(m, 15), "pago" if m < 2 else "previsto", "Parcela CNC", "Parcela financiamento máquinas", -4500)
    return L, projetos


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("saida")
    ap.add_argument("--inicio", default=dt.date.today().strftime("%Y-%m"), help="AAAA-MM do 1º mês")
    ap.add_argument("--meses", type=int, default=12)
    ap.add_argument("--saldo-inicial", type=float, default=0)
    ap.add_argument("--lancamentos")
    ap.add_argument("--exemplo", action="store_true")
    a = ap.parse_args()
    ano, mes = map(int, a.inicio.split("-"))
    inicio = dt.date(ano, mes, 1)

    linhas, projetos = [], []
    if a.lancamentos:
        linhas = ler_lancamentos(a.lancamentos)
    if a.exemplo:
        ex_l, projetos = exemplo(inicio)
        linhas += ex_l
    if len(linhas) > MAX_LINHAS - 1:
        sys.exit(f"Mais de {MAX_LINHAS - 1} lançamentos: aumente MAX_LINHAS no script.")

    wb = Workbook()
    prem = aba_premissas(wb)
    wb["Premissas"]["B4"].value = a.saldo_inicial
    lista = aba_plano(wb)
    aba_lancamentos(wb, linhas, lista)
    n_proj = aba_projetos(wb, projetos)
    linha_semcat = aba_fluxo(wb, inicio, a.meses, prem)
    aba_dre(wb, inicio, a.meses)
    aba_recebiveis(wb, n_proj)
    aba_checks(wb, linha_semcat, a.meses)
    wb.save(a.saida)
    print(f"OK: {a.saida} — {len(linhas)} lançamentos, {len(projetos)} projetos, "
          f"{a.meses} meses a partir de {inicio:%m/%Y}")


if __name__ == "__main__":
    main()
