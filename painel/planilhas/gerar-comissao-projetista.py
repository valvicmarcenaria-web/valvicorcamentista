#!/usr/bin/env python3
"""
Planilha individual de comissão do projetista — uma por pessoa, para compartilhar com ela.

    python3 gerar-comissao-projetista.py              # gera para todos da lista PESSOAS
    python3 gerar-comissao-projetista.py "Lorrane"    # gera só para uma

Duas abas: **Comissões**, onde se lança projeto a projeto, e **Painel**, com a leitura
anual e mês a mês. Uma terceira aba traz a regra escrita, porque a planilha é vista pela
pessoa que recebe — e regra de comissão que não está escrita vira discussão.

A conta: a comissão incide sobre o **líquido**, não sobre o contrato.

Os custos de venda são lançados em PERCENTUAL do contrato, não em reais — é assim que
eles são combinados (nota 6%, taxa da máquina 1,5%, RT 5%...). A planilha converte.

    custos   = valor do contrato × (nota% + taxa% + RT% + vendedor% + outros%)
    líquido  = valor do contrato − custos
    comissão = líquido × percentual do projetista (1% por padrão, editável por linha)

EDITE ESTE SCRIPT, NUNCA O .XLSX — a próxima geração sobrescreve o arquivo.

Sobre os menus suspensos: a lista precisa ser LITERAL (embutida na validação). Intervalo de
outra aba e nome definido são descartados na importação para o Google Sheets — e esta
planilha vai ser compartilhada justamente por lá.
"""
import sys
import unicodedata

from openpyxl import Workbook
from openpyxl.chart import BarChart, DoughnutChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

PESSOAS = ['Lorrane', 'Lucas']
PCT_PADRAO = 0.01
LINHAS = 120          # 12 meses × 10 projetos, o teto que o Jonathan indicou

NAVY, NAVY2, GOLD = '0E2038', '16314F', 'C2A05A'
GOLDSOFT, GOLDBG, CREME = 'D8BD80', 'F6EDD6', 'FBFAF7'
INK, MUTED, LINHA, ZEBRA = '1B2733', '6C7785', 'E8E3D8', 'F7F5EF'
OK, OKBG, ABERTO, ABERTOBG, CALC = '2F7D4F', 'E6F0E9', 'A8700F', 'FAEFDC', 'EAF1F8'

_f = Side(style='thin', color='D6DBE2')
BORDA = Border(left=_f, right=_f, top=_f, bottom=_f)
DIN = 'R$ #,##0.00'
PCT = '0.0%'
DATA = 'DD/MM/YYYY'

MESES = ['Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun', 'Jul', 'Ago', 'Set', 'Out', 'Nov', 'Dez']
MES_EXT = ['Janeiro', 'Fevereiro', 'Março', 'Abril', 'Maio', 'Junho',
           'Julho', 'Agosto', 'Setembro', 'Outubro', 'Novembro', 'Dezembro']
STATUS = ['Em aberto', 'Pago']

# (título, largura, tipo)
#   'e' preenche        'p' preenche em percentual do contrato
#   'c' calcula         'v' calcula e fica em destaque (a comissão do projetista)
COLS = [
    ('Mês', 8, 'e'), ('Cliente', 24, 'e'), ('Ambiente / projeto', 26, 'e'),
    ('Valor do contrato', 16, 'e'),
    ('Nota fiscal', 11, 'p'), ('Taxa de máquina', 12, 'p'), ('RT', 11, 'p'),
    ('Comissão do vendedor', 13, 'p'), ('Outros custos', 11, 'p'),
    ('Total de custos', 15, 'c'), ('Valor líquido', 16, 'c'),
    ('%', 7, 'e'), ('Comissão a receber', 18, 'v'),
    ('Status', 13, 'e'), ('Pagamento em', 13, 'e'),
    ('Conferir', 26, 'c'),
]
CUSTOS = 'E'            # primeira coluna de custo
CUSTOS_FIM = 'I'        # última coluna de custo
TETO_CUSTO = 0.40       # soma de custos acima disto acende a coluna Conferir.
                        # Na prática os projetos ficam entre 15% e 22%, então 40% dá folga
                        # e ainda pega o erro clássico: 60% digitado no lugar de 6%.
PRIMEIRA, ULTIMA = 4, 3 + LINHAS


def literal(itens):
    for it in itens:
        assert ',' not in it and '"' not in it, f'item quebra a lista literal: {it!r}'
    lit = '"' + ','.join(itens) + '"'
    assert len(lit) <= 255, f'lista literal com {len(lit)} caracteres (máx. 255)'
    return lit


def menu(ws, itens, intervalo):
    dv = DataValidation(type='list', formula1=literal(itens),
                        allow_blank=True, showDropDown=False)   # False = MOSTRA a setinha
    dv.errorTitle, dv.error = 'Valor fora da lista', 'Escolha um valor da lista.'
    ws.add_data_validation(dv)
    dv.add(intervalo)


def faixa(ws, linha, texto, sub, ate):
    ws.merge_cells(start_row=linha, start_column=1, end_row=linha, end_column=ate)
    c = ws.cell(row=linha, column=1, value=texto)
    c.fill = PatternFill('solid', fgColor=NAVY)
    c.font = Font(name='Calibri', size=14, bold=True, color=GOLDSOFT)
    c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[linha].height = 28
    ws.merge_cells(start_row=linha + 1, start_column=1, end_row=linha + 1, end_column=ate)
    c = ws.cell(row=linha + 1, column=1, value=sub)
    c.fill = PatternFill('solid', fgColor=GOLDBG)
    c.font = Font(name='Calibri', size=9.5, italic=True, color=NAVY2)
    c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[linha + 1].height = 19


def aba_comissoes(wb, nome):
    ws = wb.active
    ws.title = 'Comissões'
    ws.sheet_properties.tabColor = GOLD
    faixa(ws, 1, f'VALVIC MARCENARIA   ·   COMISSÕES   ·   {nome.upper()}',
          'Os custos de venda entram em PERCENTUAL do contrato. A comissão incide sobre o '
          'valor líquido. Preencha as colunas claras; as azuis se calculam sozinhas.', len(COLS))

    for i, (titulo, larg, tipo) in enumerate(COLS, start=1):
        c = ws.cell(row=3, column=i, value=titulo)
        c.fill = PatternFill('solid', fgColor=GOLD if tipo == 'v' else
                             (NAVY2 if tipo == 'c' else NAVY))
        c.font = Font(name='Calibri', size=9.5, bold=True,
                      color=NAVY if tipo == 'v' else GOLDSOFT)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BORDA
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.row_dimensions[3].height = 32
    ws.freeze_panes = 'D4'

    din = {4, 10, 11, 13}       # colunas em reais
    pct = {5, 6, 7, 8, 9, 12}   # colunas em percentual
    for r in range(PRIMEIRA, ULTIMA + 1):
        par = (r - PRIMEIRA) % 2
        for i, (_, _, tipo) in enumerate(COLS, start=1):
            c = ws.cell(row=r, column=i)
            c.border = BORDA
            if tipo == 'v':
                c.fill = PatternFill('solid', fgColor=GOLDBG)
                c.font = Font(name='Calibri', size=11, bold=True, color=NAVY)
            elif tipo == 'c':
                c.fill = PatternFill('solid', fgColor=CALC)
                c.font = Font(name='Calibri', size=10, color=NAVY2, italic=True)
            else:
                c.fill = PatternFill('solid', fgColor=ZEBRA if par else 'FFFFFF')
                c.font = Font(name='Calibri', size=10, color=INK)
            if i in din:
                c.number_format = DIN
            elif i in pct:
                c.number_format = PCT
        ws.cell(row=r, column=15).number_format = DATA
        soma = f'SUM(${CUSTOS}{r}:${CUSTOS_FIM}{r})'
        # ── as contas ──
        ws.cell(row=r, column=10, value=f'=IF($D{r}="","",ROUND($D{r}*{soma},2))')
        ws.cell(row=r, column=11, value=f'=IF($D{r}="","",$D{r}-$J{r})')
        ws.cell(row=r, column=12, value=f'=IF($D{r}="","",{PCT_PADRAO})')
        ws.cell(row=r, column=13, value=f'=IF($D{r}="","",ROUND($K{r}*$L{r},2))')
        ws.cell(row=r, column=16, value=(
            f'=IF($D{r}="","",'
            f'IF({soma}>={TETO_CUSTO},"custos acima de {TETO_CUSTO:.0%} do contrato",'
            f'IF(AND($N{r}="Pago",$O{r}=""),"marcado como pago, sem a data",'
            f'IF($A{r}="","falta o mês",""))))'))
        ws.cell(row=r, column=16).font = Font(name='Calibri', size=9.5, color='B0413F', bold=True)

    menu(ws, MESES, f'A{PRIMEIRA}:A{ULTIMA}')
    menu(ws, STATUS, f'N{PRIMEIRA}:N{ULTIMA}')

    verde = PatternFill('solid', fgColor=OKBG)
    ws.conditional_formatting.add(f'N{PRIMEIRA}:N{ULTIMA}',
        FormulaRule(formula=[f'$N{PRIMEIRA}="Pago"'], fill=verde, font=Font(color=OK, bold=True)))
    ws.conditional_formatting.add(f'N{PRIMEIRA}:N{ULTIMA}',
        FormulaRule(formula=[f'$N{PRIMEIRA}="Em aberto"'],
                    fill=PatternFill('solid', fgColor=ABERTOBG), font=Font(color=ABERTO, bold=True)))
    ws.conditional_formatting.add(f'P{PRIMEIRA}:P{ULTIMA}',
        FormulaRule(formula=[f'$P{PRIMEIRA}<>""'], fill=PatternFill('solid', fgColor='FBE3E2')))

    # filtro no cabecalho: permite isolar um mes ou um status sem mexer em nada
    ws.auto_filter.ref = f'A3:P{ULTIMA}'

    # impressao: paisagem, uma pagina de largura, cabecalho repetido a cada folha
    ws.page_setup.orientation = 'landscape'
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '1:3'
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.3
    ws.page_margins.top = ws.page_margins.bottom = 0.4
    return ws


def aba_painel(wb, nome):
    ws = wb.create_sheet('Painel')
    ws.sheet_properties.tabColor = NAVY2
    for col, larg in zip('ABCDEFGHI', [4, 24, 15, 15, 15, 15, 15, 15, 15]):
        ws.column_dimensions[col].width = larg
    faixa(ws, 1, f'PAINEL   ·   {nome.upper()}',
          'Tudo se calcula a partir da aba Comissões. Nada se digita aqui.', 9)

    C = "Comissões"
    def band(linha, texto):
        ws.merge_cells(start_row=linha, start_column=2, end_row=linha, end_column=9)
        c = ws.cell(row=linha, column=2, value=texto)
        c.fill = PatternFill('solid', fgColor=GOLDBG)
        c.font = Font(name='Calibri', size=10, bold=True, color=NAVY2)
        c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        ws.row_dimensions[linha].height = 20

    # ── quatro indicadores do ano, lado a lado ──
    band(4, 'NO ANO')
    kpis = [
        ('Projetos lançados', f'=COUNTIF(\'{C}\'!$D:$D,">0")', '0', NAVY),
        ('Comissão total', f"=SUM('{C}'!$M:$M)", DIN, NAVY),
        ('Já paga', f"=SUMIF('{C}'!$N:$N,\"Pago\",'{C}'!$M:$M)", DIN, OK),
        ('Em aberto', f"=SUMIF('{C}'!$N:$N,\"Em aberto\",'{C}'!$M:$M)", DIN, ABERTO),
    ]
    for i, (rot, formula, fmt, cor) in enumerate(kpis):
        col = 2 + i * 2
        r = ws.cell(row=5, column=col, value=rot)
        r.font = Font(name='Calibri', size=9, bold=True, color=MUTED)
        ws.merge_cells(start_row=5, start_column=col, end_row=5, end_column=col + 1)
        v = ws.cell(row=6, column=col, value=formula)
        v.number_format = fmt
        v.font = Font(name='Calibri', size=16, bold=True, color=cor)
        v.alignment = Alignment(horizontal='left', vertical='center')
        ws.merge_cells(start_row=6, start_column=col, end_row=6, end_column=col + 1)
    ws.row_dimensions[6].height = 26

    # ── mês a mês ──
    band(8, 'MÊS A MÊS')
    cab = ['Mês', 'Projetos', 'Comissão', 'Paga', 'Em aberto']
    ULT = 2 + len(cab) - 1                      # última coluna da tabela (F)
    for i, t in enumerate(cab):
        c = ws.cell(row=9, column=2 + i, value=t)
        c.fill = PatternFill('solid', fgColor=NAVY)
        c.font = Font(name='Calibri', size=9.5, bold=True, color=GOLDSOFT)
        c.alignment = Alignment(horizontal='center', vertical='center')
        c.border = BORDA
    ws.row_dimensions[9].height = 22

    P1, P2 = 10, 21                             # primeiro e último mês
    for i, (m, ext) in enumerate(zip(MESES, MES_EXT)):
        r = P1 + i
        ws.cell(row=r, column=2, value=ext).font = Font(name='Calibri', size=10, bold=True, color=NAVY2)
        ws.cell(row=r, column=3, value=f"=COUNTIFS('{C}'!$A:$A,\"{m}\",'{C}'!$D:$D,\">0\")")
        ws.cell(row=r, column=4, value=f"=SUMIF('{C}'!$A:$A,\"{m}\",'{C}'!$M:$M)")
        ws.cell(row=r, column=5, value=f"=SUMIFS('{C}'!$M:$M,'{C}'!$A:$A,\"{m}\",'{C}'!$N:$N,\"Pago\")")
        ws.cell(row=r, column=6, value=f"=SUMIFS('{C}'!$M:$M,'{C}'!$A:$A,\"{m}\",'{C}'!$N:$N,\"Em aberto\")")
        for c in range(2, ULT + 1):
            cel = ws.cell(row=r, column=c)
            cel.border = BORDA
            cel.fill = PatternFill('solid', fgColor=ZEBRA if i % 2 else 'FFFFFF')
            if c >= 4:
                cel.number_format = DIN
            if c == 3:
                cel.alignment = Alignment(horizontal='center')
            if c != 2:
                cel.font = Font(name='Calibri', size=10, color=NAVY2)

    TOT = P2 + 1                                # linha do total do ano (22)
    ws.cell(row=TOT, column=2, value='Total do ano').font = Font(
        name='Calibri', size=10, bold=True, color=GOLDSOFT)
    for c in range(3, ULT + 1):
        L = get_column_letter(c)
        cel = ws.cell(row=TOT, column=c, value=f'=SUM({L}{P1}:{L}{P2})')
        cel.number_format = '0' if c == 3 else DIN
        cel.font = Font(name='Calibri', size=10, bold=True, color=GOLDSOFT)
    for c in range(2, ULT + 1):
        cel = ws.cell(row=TOT, column=c)
        cel.fill = PatternFill('solid', fgColor=NAVY)
        cel.border = BORDA
        if c == 3:
            cel.alignment = Alignment(horizontal='center')

    # ── gráficos: o ano mês a mês, e a divisão entre pago e em aberto ──
    band(24, 'GRÁFICOS')

    barras = BarChart()
    barras.type = 'col'
    barras.title = 'Comissão por mês'
    barras.y_axis.title = None
    barras.x_axis.title = None
    barras.legend = None
    barras.height, barras.width = 8.6, 14.5
    barras.add_data(Reference(ws, min_col=4, min_row=9, max_row=P2), titles_from_data=True)
    barras.set_categories(Reference(ws, min_col=2, min_row=P1, max_row=P2))
    barras.series[0].graphicalProperties.solidFill = GOLD
    barras.series[0].graphicalProperties.line.solidFill = GOLD
    ws.add_chart(barras, 'B26')

    rosca = DoughnutChart(holeSize=52)
    rosca.title = 'Comissão paga e em aberto'
    rosca.height, rosca.width = 8.6, 9.5
    rosca.add_data(Reference(ws, min_col=5, max_col=6, min_row=TOT, max_row=TOT), from_rows=True)
    rosca.set_categories(Reference(ws, min_col=5, max_col=6, min_row=9, max_row=9))
    rosca.dataLabels = DataLabelList()
    rosca.dataLabels.showVal = False
    rosca.dataLabels.showPercent = True
    ws.add_chart(rosca, 'F26')

    return ws


def aba_regra(wb, nome):
    ws = wb.create_sheet('Como funciona')
    ws.sheet_properties.tabColor = MUTED
    for col, larg in zip('ABC', [4, 30, 86]):
        ws.column_dimensions[col].width = larg
    faixa(ws, 1, 'COMO A COMISSÃO É CALCULADA',
          'A mesma regra para todos os projetos. Qualquer dúvida se resolve olhando esta aba.', 3)
    blocos = [
        ('A conta', 'A comissão incide sobre o VALOR LÍQUIDO do projeto, nunca sobre o valor de contrato. '
                    'Líquido = valor do contrato − custos de venda. Os custos de venda são a soma dos '
                    'percentuais lançados (nota fiscal, taxa de máquina, RT, comissão do vendedor e outros) '
                    'aplicada sobre o contrato.'),
        ('O percentual', f'{PCT_PADRAO:.0%} sobre o líquido. A coluna "%" existe em cada linha porque um projeto '
                         'pode ter percentual diferente — quando houver, o combinado é anotado ali antes do lançamento.'),
        ('Custos de venda', 'São os custos que a venda carrega e que não ficam com a empresa: imposto da nota '
                            'fiscal, taxa da máquina de cartão, RT da arquiteta ou decoradora, a comissão do '
                            'vendedor e eventuais outros. Cada um tem a sua coluna e é lançado em PERCENTUAL '
                            'do contrato — é como eles são combinados. A planilha converte para reais.'),
        ('Status', 'EM ABERTO — a comissão está na fila, a pagar. PAGO — pago, com a data preenchida. '
                   'Só entra nesta planilha o projeto que já foi direcionado para ela.'),
        ('Quem preenche', 'A Valvic lança e atualiza. A planilha é compartilhada com o projetista para '
                          'acompanhamento — conferir é bem-vindo; qualquer divergência se fala antes do pagamento.'),
        ('A coluna Conferir', 'Acende sozinha quando falta o mês, quando o projeto está marcado como pago sem '
                              f'data, ou quando os percentuais de custo somam mais de {TETO_CUSTO:.0%} do '
                              'contrato. É uma rede contra erro de digitação: um 60% no lugar de 6%. Projeto '
                              'que realmente passe disso existe — aí é só conferir e seguir.'),
    ]
    r = 4
    for titulo, texto in blocos:
        c = ws.cell(row=r, column=2, value=titulo)
        c.font = Font(name='Calibri', size=11, bold=True, color=NAVY)
        c.alignment = Alignment(vertical='top')
        t = ws.cell(row=r, column=3, value=texto)
        t.font = Font(name='Calibri', size=10.5, color=INK)
        t.alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[r].height = 46
        r += 1
    return ws


def gerar(nome):
    wb = Workbook()
    aba_comissoes(wb, nome)
    aba_painel(wb, nome)
    aba_regra(wb, nome)
    base = ''.join(ch for ch in unicodedata.normalize('NFD', nome)
                   if unicodedata.category(ch) != 'Mn')
    arq = f'Valvic_Comissao_{base.replace(" ", "_")}.xlsx'
    wb.save(arq)
    print(f'planilha gerada: {arq}')
    return arq


if __name__ == '__main__':
    alvos = sys.argv[1:] or PESSOAS
    for n in alvos:
        gerar(n)
