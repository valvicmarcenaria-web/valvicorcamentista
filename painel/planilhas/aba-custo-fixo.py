#!/usr/bin/env python3
"""Aba "Custo Fixo" da Controle de Pagamentos — o custo vigente da operação.

    python3 aba-custo-fixo.py ENTRADA.xlsx SAIDA.xlsx

Vem da planilha de referência de ago/2025, enxugada. O que mudou e por quê:

  · a referência gastava sete linhas por pessoa (salário, VT, cesta, exames,
    bônus, encargos, total). Aqui cada pessoa é UMA linha com colunas — dá para
    ver a equipe inteira de uma vez, e acrescentar alguém é digitar uma linha;
  · os encargos deixam de ser digitados: saem do vínculo. CLT calcula, PJ e
    sócio não têm. Era onde a referência mais errava, porque o 33% estava
    repetido à mão em cada bloco;
  · a referência misturava custo fixo com financiamento no mesmo total. Aqui
    são três blocos separados, como pedido: FOLHA, ESTRUTURA e INVESTIMENTOS;
  · o patrimônio (bens, veículos, máquinas) saiu: não é custo, é ativo. Ele já
    tem lugar próprio na Valvic_Controle_Patrimonio.

A aba é lida pelo painel do Walton: os três subtotais e o custo operacional.
"""
import datetime
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

NAVY, NAVY2, GOLD = '0E2038', '16314F', 'C2A05A'
GOLDS, GOLDBG, INPUT, CALC = 'D8BD80', 'F6EDD6', 'FBF7EC', 'EEF1F4'
INK, MUTED, LINE, ZEBRA = '1B2733', '6C7785', 'D6DBE2', 'F7F5EF'
MOEDA, PCT = 'R$ #,##0.00', '0%'
F = 'Arial'
BORDA = Border(left=Side('thin', color=LINE), right=Side('thin', color=LINE),
               top=Side('thin', color=LINE), bottom=Side('thin', color=LINE))

ENCARGO_CLT = 0.33      # férias, 13º, FGTS e multa — o mesmo índice da referência
VINCULOS = ['CLT', 'PJ', 'Sócio']

# ── a folha de hoje ─────────────────────────────────────────────────────
# nome, função, vínculo, fixo, VT, VR/cesta
FOLHA = [
    ('Jonathan Oliveira', 'Diretor Comercial & Experiência', 'Sócio', 10000, 0, 0),
    ('Paulo Henrique', 'Diretor de Operações & Tecnologia', 'Sócio', 10000, 0, 0),
    ('Ronald', 'Marceneiro', 'PJ', 6500, 550, 0),
    ('Lorrane', 'Projetista', 'PJ', 4800, 0, 0),
    ('Deivison', 'Coordenador de Produção', 'PJ', 4000, 0, 0),
    ('Bruna', 'Arquiteta · projeto e programação', 'PJ', 4000, 0, 0),
    ('Samuel', 'Marceneiro Sênior', 'PJ', 3900, 0, 0),
    ('Cesar', 'Marceneiro', 'PJ', 3900, 0, 0),
    ('Joelson', 'Operador CNC & máquinas', 'PJ', 3800, 0, 0),
    ('Jhon', 'Marceneiro', 'PJ', 3300, 0, 0),
    ('Lucas', 'Projetista', 'CLT', 3500, 520, 0),
    ('Davi', 'Marceneiro Júnior', 'CLT', 2200, 450, 300),
    ('Karla', 'Administrativo', 'CLT', 2000, 550, 500),
    ('Jomar', 'Marceneiro', 'CLT', 2000, 400, 300),
    ('Bruno', 'Auxiliar de Marceneiro', 'CLT', 2000, 300, 300),
    ('Wallace', 'Trainee', 'CLT', 1750, 400, 300),
    ('Douglas', 'Trainee', 'CLT', 1620, 300, 300),
    ('Alex', 'Ajudante Geral', 'CLT', 1620, 300, 300),
]
VAGAS_FOLHA = 8          # linhas livres para quem entrar

# ── estrutura · o que a empresa paga mesmo sem vender ───────────────────
ESTRUTURA = [
    ('Aluguel do barracão e IPTU', 5000), ('Marketing, carros e uniformes', 4000),
    ('Combustível', 3000), ('Depreciação', 2500), ('Luz do barracão', 1500),
    ('Contabilidade', 1000), ('Alvarás, taxas e e-Social', 1000),
    ('IPVA de caminhão e carros', 800), ('Faxina do barracão', 500),
    ('Vigilância eletrônica', 400), ('Desenvolvimento e cursos', 400),
    ('Sistema Calcme', 320), ('Produtos de limpeza', 300), ('Água', 200),
    ('Manutenção periódica das máquinas', 200), ('Materiais de escritório', 150),
    ('Seguro de veículos', 150), ('Internet da fábrica', 130),
    ('Seguro de colaboradores', 130), ('Coleta de resíduos', 100),
    ('Tarifas bancárias', 100), ('Extintor', 100),
]
VAGAS_ESTRUTURA = 6

# ── investimentos · financiamento e empréstimo ──────────────────────────
INVESTIMENTOS = [
    ('Financiamento de máquinas', 6000), ('Empréstimo Paulo', 5000),
    ('Financiamento de carros', 3000), ('Consórcio Itaú', 1100),
    ('Substituição tributária', 150),
]
VAGAS_INVEST = 5


def faixa(ws, r, titulo, sub, ate):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ate)
    c = ws.cell(r, 1, titulo)
    c.fill = PatternFill('solid', fgColor=NAVY)
    c.font = Font(name=F, size=13, bold=True, color=GOLDS)
    c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[r].height = 26
    ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=ate)
    c = ws.cell(r + 1, 1, sub)
    c.fill = PatternFill('solid', fgColor=GOLDBG)
    c.font = Font(name=F, size=9, italic=True, color=NAVY2)
    c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[r + 1].height = 18


def secao(ws, r, texto, ate):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ate)
    c = ws.cell(r, 1, '  ' + texto)
    c.fill = PatternFill('solid', fgColor=NAVY2)
    c.font = Font(name=F, size=9.5, bold=True, color=GOLDS)
    c.alignment = Alignment(horizontal='left', vertical='center')
    ws.row_dimensions[r].height = 20


def cab(ws, r, titulos):
    for i, t in enumerate(titulos, start=1):
        c = ws.cell(r, i, t)
        c.fill = PatternFill('solid', fgColor='41505D')
        c.font = Font(name=F, size=8, bold=True, color='FFFFFF')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BORDA
    ws.row_dimensions[r].height = 24


def total(ws, r, rotulo, col_valores, ini, fim, ate, destaque=False):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=col_valores - 1)
    c = ws.cell(r, 1, rotulo)
    c.font = Font(name=F, size=10 if destaque else 9.5, bold=True, color=GOLDS)
    c.alignment = Alignment(horizontal='right', vertical='center', indent=1)
    L = get_column_letter(col_valores)
    v = ws.cell(r, col_valores, f'=ROUND(SUM({L}{ini}:{L}{fim}),2)')
    v.number_format = MOEDA
    v.font = Font(name=F, size=11 if destaque else 10, bold=True, color=GOLDS)
    for c0 in range(1, ate + 1):
        ws.cell(r, c0).fill = PatternFill('solid', fgColor=NAVY if destaque else NAVY2)
        ws.cell(r, c0).border = BORDA
    ws.row_dimensions[r].height = 22 if destaque else 20
    return r


def main(entrada, saida):
    wb = openpyxl.load_workbook(entrada)
    # a aba de meses dava o mesmo número em outro lugar — duas fontes do mesmo
    # dado sempre divergem, então ela sai e fica só esta
    for velha in ('Custo da Empresa', 'Custo Fixo'):
        if velha in wb.sheetnames:
            del wb[velha]

    ws = wb.create_sheet('Custo Fixo')
    ws.sheet_properties.tabColor = GOLD
    ws.sheet_view.showGridLines = False
    NC = 9
    for col, larg in zip('ABCDEFGHI', (26, 30, 11, 14, 11, 11, 13, 14, 26)):
        ws.column_dimensions[col].width = larg

    hoje = datetime.date.today()
    faixa(ws, 1, 'CUSTO FIXO DA OPERAÇÃO',
          f'Atualizado em {hoje.strftime("%d/%m/%Y")}  ·  preencha e o resto se calcula  ·  '
          f'os encargos saem do vínculo, não se digita', NC)

    # ── resumo, no topo, porque é o que o painel lê ──
    r = 4
    secao(ws, r, 'O QUE A OPERAÇÃO CUSTA POR MÊS', NC)
    r += 1
    resumo_linha = r
    for i, rot in enumerate(('Folha de pagamento', 'Estrutura', 'Investimentos')):
        c = ws.cell(r, 1 + i * 3, rot)
        c.font = Font(name=F, size=8.5, bold=True, color=MUTED)
        ws.merge_cells(start_row=r, start_column=1 + i * 3, end_row=r, end_column=3 + i * 3)
        v = ws.cell(r + 1, 1 + i * 3)
        v.number_format = MOEDA
        v.font = Font(name=F, size=15, bold=True, color=NAVY)
        v.fill = PatternFill('solid', fgColor=GOLDBG)
        v.alignment = Alignment(horizontal='center', vertical='center')
        ws.merge_cells(start_row=r + 1, start_column=1 + i * 3, end_row=r + 1, end_column=3 + i * 3)
    ws.row_dimensions[r + 1].height = 28
    r += 3

    # ── 1 · folha ──
    secao(ws, r, '1 · FOLHA DE PAGAMENTO', NC); r += 1
    cab(ws, r, ['Nome', 'Função', 'Vínculo', 'Fixo mensal', 'VT', 'VR / cesta',
                'Encargos', 'Total', 'Observação']); r += 1
    f0 = r
    for i in range(len(FOLHA) + VAGAS_FOLHA):
        dados = FOLHA[i] if i < len(FOLHA) else None
        for c in range(1, NC + 1):
            cel = ws.cell(r, c)
            cel.border = BORDA
            cel.font = Font(name=F, size=9, color=INK)
            cel.fill = PatternFill('solid', fgColor=ZEBRA if i % 2 else 'FFFFFF')
        if dados:
            nome, func, vinc, fixo, vt, vr = dados
            ws.cell(r, 1, nome).font = Font(name=F, size=9, bold=True, color=NAVY)
            ws.cell(r, 2, func)
            ws.cell(r, 3, vinc)
            for c, v in ((4, fixo), (5, vt), (6, vr)):
                ws.cell(r, c, v or None)
        ws.cell(r, 3).alignment = Alignment(horizontal='center')
        for c in (4, 5, 6):
            ws.cell(r, c).number_format = MOEDA
        enc = ws.cell(r, 7, f'=IF($A{r}="","",IF($C{r}="CLT",'
                            f'ROUND($D{r}*{ENCARGO_CLT},2),0))')
        enc.number_format = MOEDA
        enc.fill = PatternFill('solid', fgColor=CALC)
        enc.font = Font(name=F, size=9, color=NAVY2, italic=True)
        tot = ws.cell(r, 8, f'=IF($A{r}="","",ROUND(SUM($D{r}:$G{r}),2))')
        tot.number_format = MOEDA
        tot.fill = PatternFill('solid', fgColor=CALC)
        tot.font = Font(name=F, size=9.5, bold=True, color=NAVY)
        r += 1
    ff = r - 1
    dv = DataValidation(type='list', formula1='"' + ','.join(VINCULOS) + '"',
                        allow_blank=True, showDropDown=False)
    ws.add_data_validation(dv)
    dv.add(f'C{f0}:C{ff}')
    lf = total(ws, r, 'TOTAL DA FOLHA  →', 8, f0, ff, NC); r += 2

    # ── 2 · estrutura ──
    secao(ws, r, '2 · ESTRUTURA', NC); r += 1
    cab(ws, r, ['Item', '', '', '', '', '', '', 'Valor mensal', 'Observação']); r += 1
    e0 = r
    for i in range(len(ESTRUTURA) + VAGAS_ESTRUTURA):
        for c in range(1, NC + 1):
            cel = ws.cell(r, c)
            cel.border = BORDA
            cel.font = Font(name=F, size=9, color=INK)
            cel.fill = PatternFill('solid', fgColor=ZEBRA if i % 2 else 'FFFFFF')
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        if i < len(ESTRUTURA):
            ws.cell(r, 1, ESTRUTURA[i][0])
            ws.cell(r, 8, ESTRUTURA[i][1])
        ws.cell(r, 8).number_format = MOEDA
        r += 1
    ef = r - 1
    le = total(ws, r, 'TOTAL DA ESTRUTURA  →', 8, e0, ef, NC); r += 2

    # ── 3 · investimentos ──
    secao(ws, r, '3 · INVESTIMENTOS', NC); r += 1
    cab(ws, r, ['Item', '', '', '', '', '', '', 'Valor mensal', 'Observação']); r += 1
    i0 = r
    for i in range(len(INVESTIMENTOS) + VAGAS_INVEST):
        for c in range(1, NC + 1):
            cel = ws.cell(r, c)
            cel.border = BORDA
            cel.font = Font(name=F, size=9, color=INK)
            cel.fill = PatternFill('solid', fgColor=ZEBRA if i % 2 else 'FFFFFF')
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        if i < len(INVESTIMENTOS):
            ws.cell(r, 1, INVESTIMENTOS[i][0])
            ws.cell(r, 8, INVESTIMENTOS[i][1])
        ws.cell(r, 8).number_format = MOEDA
        r += 1
    if_ = r - 1
    li = total(ws, r, 'TOTAL DOS INVESTIMENTOS  →', 8, i0, if_, NC); r += 2

    # ── fecha os totais do topo ──
    ws.cell(resumo_linha + 1, 1, f'=$H${lf}')
    ws.cell(resumo_linha + 1, 4, f'=$H${le}')
    ws.cell(resumo_linha + 1, 7, f'=$H${li}')

    # ── fechamento ──
    for rot, formula, dest in (
            ('CUSTO FIXO OPERACIONAL   ·   folha + estrutura', f'=ROUND($H${lf}+$H${le},2)', True),
            ('COM INVESTIMENTOS', f'=ROUND($H${lf}+$H${le}+$H${li},2)', False)):
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=7)
        c = ws.cell(r, 1, rot)
        c.font = Font(name=F, size=11 if dest else 10, bold=True, color=GOLDS)
        c.alignment = Alignment(horizontal='right', vertical='center', indent=1)
        v = ws.cell(r, 8, formula)
        v.number_format = MOEDA
        v.font = Font(name=F, size=13 if dest else 11, bold=True, color=GOLDS)
        for c0 in range(1, NC + 1):
            ws.cell(r, c0).fill = PatternFill('solid', fgColor=NAVY if dest else NAVY2)
            ws.cell(r, c0).border = BORDA
        ws.row_dimensions[r].height = 26 if dest else 22
        r += 1

    r += 1
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=NC)
    n = ws.cell(r, 1,
                f'  Encargos: {ENCARGO_CLT:.0%} sobre o fixo, só para CLT — férias, 13º, FGTS e '
                f'multa. Confirme o índice com a contabilidade. PJ e sócio não têm encargo, o '
                f'custo é o valor do contrato.   ·   Para incluir alguém, use uma linha em branco '
                f'da folha: nome, função, vínculo e os valores. O resto se calcula.   ·   '
                f'Patrimônio não entra aqui: bem é ativo, não custo — fica na Controle de '
                f'Patrimônio.')
    n.font = Font(name=F, size=8.5, color='41505D')
    n.fill = PatternFill('solid', fgColor=GOLDBG)
    n.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    ws.row_dimensions[r].height = 42

    ws.freeze_panes = 'A4'
    wb.save(saida)
    print('aba "Custo Fixo" criada')
    print(f'  folha ......... linhas {f0}-{ff}  ({len(FOLHA)} pessoas + {VAGAS_FOLHA} vagas)')
    print(f'  estrutura ..... linhas {e0}-{ef}  ({len(ESTRUTURA)} itens + {VAGAS_ESTRUTURA} vagas)')
    print(f'  investimentos . linhas {i0}-{if_}  ({len(INVESTIMENTOS)} itens + {VAGAS_INVEST} vagas)')
    print(f'  abas no arquivo: {len(wb.sheetnames)} — {", ".join(wb.sheetnames)}')
    print(f'\nOK → {saida}')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
