#!/usr/bin/env python3
"""
Gera a planilha de gestão mensal conjunta com o Walton.

    python3 gerar-gestao-mensal-walton.py

Saída: Valvic_Gestao_Mensal_2026.xlsx

EDITE ESTE SCRIPT, NUNCA O .XLSX — a próxima geração sobrescreve o arquivo.

O desenho está na lâmina painel/esquema-planilha-walton.html. Em uma frase: as cinco
métricas que o Walton acompanha são cinco eventos com data e fonte próprias, então há
cinco abas de lançamento e treze lâminas que só calculam.

Duas decisões que explicam quase todas as fórmulas daqui:

1. A unidade de lançamento é o AMBIENTE, não o projeto. Projeto nunca fica "fabricado";
   fica fabricado aos pedaços. Ambiente fica pronto ou não fica — e isso tira a estimativa
   do meio do número.

2. Material entra pelo USO, não pela data da compra. A compra é lançada inteira e a
   planilha a rateia entre os ambientes do projeto, na proporção do valor de cada um; o
   material de cada ambiente conta no mês em que aquele ambiente foi fabricado. Por isso
   "material ÷ fabricado" é uma margem de verdade, e não uma coincidência de calendário.

Sobre os menus suspensos: a lista precisa ser LITERAL (embutida na validação). Intervalo de
outra aba e nome definido são descartados na importação para o Google Sheets — foi assim que
a planilha de custos saiu sem menu nenhum. Limite de 255 caracteres, sem vírgula nos itens;
a função literal() confere isso e falha alto.

Para o campo Projeto não existe menu: a lista de projetos passaria de 255 caracteres e
cresce ao longo do ano. No lugar dele há uma coluna de aviso que acende quando o nome
digitado não está cadastrado em Projetos, e a conferência Σ ambientes × valor do contrato,
que pega o mesmo erro por outro caminho.
"""
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ARQUIVO = 'Valvic_Gestao_Mensal_2026.xlsx'
ANO = 2026

# ── paleta da casa ────────────────────────────────────────────────────────────
NAVY, NAVY2, GOLD = '0E2038', '16314F', 'C2A05A'
GOLDSOFT, GOLDBG = 'D8BD80', 'F6EDD6'
INK, MUTED, CINZA, ZEBRA = '1B2733', '6C7785', 'F2F4F7', 'F7F9FB'
OK, RED, BLUE, ROXO = '2F7D4F', 'B0413F', '2F5D8C', '6B4E8C'
CALC = 'EAF1F8'      # fundo das colunas calculadas
ENTRADA = 'FFFDF6'   # fundo das colunas que se preenche

_f = Side(style='thin', color='D6DBE2')
BORDA = Border(left=_f, right=_f, top=_f, bottom=_f)
DINHEIRO = 'R$ #,##0'
PCT = '0.0%'

MESES = ['Jan', 'Fev', 'Mar', 'Abr', 'Mai', 'Jun', 'Jul', 'Ago', 'Set', 'Out', 'Nov', 'Dez']
MES_EXTENSO = {'Jan': 'Janeiro', 'Fev': 'Fevereiro', 'Mar': 'Março', 'Abr': 'Abril',
               'Mai': 'Maio', 'Jun': 'Junho', 'Jul': 'Julho', 'Ago': 'Agosto',
               'Set': 'Setembro', 'Out': 'Outubro', 'Nov': 'Novembro', 'Dez': 'Dezembro'}

STATUS = ['Vendido', 'Em produção', 'Em montagem', 'Entregue', 'Pausado', 'Cancelado']
VENDEDORES = ['Jonathan', 'Paulo', 'Parceria', 'Indicação', 'Instagram', 'Outro']
CATEGORIAS = ['Chapa', 'Fita de borda', 'Ferragem', 'Iluminação', 'Vidro e espelho',
              'Serralheria', 'Marmoraria', 'Eletro', 'Insumo de fábrica',
              'Terceirizado', 'Frete', 'Outros']

LINHAS_PROJETOS = 40
LINHAS_AMBIENTES = 250
LINHAS_FATURAMENTO = 150
LINHAS_COMPRAS = 300


# ── ajudantes ─────────────────────────────────────────────────────────────────
def literal(itens):
    """Lista embutida na validação — o único formato que sobrevive em todo programa."""
    for it in itens:
        assert ',' not in it, f'item com vírgula quebra a lista literal: {it!r}'
        assert '"' not in it, f'item com aspas quebra a lista literal: {it!r}'
    lit = '"' + ','.join(itens) + '"'
    assert len(lit) <= 255, f'lista literal com {len(lit)} caracteres (máx. 255): {lit[:60]}…'
    return lit


def menu(ws, itens, intervalo):
    dv = DataValidation(type='list', formula1=literal(itens),
                        allow_blank=True, showDropDown=False)  # False = MOSTRA a setinha
    dv.errorTitle, dv.error = 'Valor fora da lista', 'Escolha um valor da lista.'
    ws.add_data_validation(dv)
    dv.add(intervalo)


def titulo(ws, texto, sub, ate):
    """Faixa de identificação no topo da aba."""
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ate)
    c = ws.cell(row=1, column=1, value=f'VALVIC MARCENARIA   ·   {texto}')
    c.fill = PatternFill('solid', fgColor=NAVY)
    c.font = Font(name='Calibri', size=13, bold=True, color=GOLDSOFT)
    c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[1].height = 26

    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ate)
    c = ws.cell(row=2, column=2 - 1, value=sub)
    c.fill = PatternFill('solid', fgColor=GOLDBG)
    c.font = Font(name='Calibri', size=9, italic=True, color=NAVY2)
    c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[2].height = 18


def cabecalho(ws, colunas, linha=3):
    """colunas = [(título, largura, 'entrada'|'calc'), ...]"""
    for i, (nome, larg, tipo) in enumerate(colunas, start=1):
        c = ws.cell(row=linha, column=i, value=nome)
        c.fill = PatternFill('solid', fgColor=NAVY if tipo == 'entrada' else BLUE)
        c.font = Font(name='Calibri', size=9, bold=True, color=GOLDSOFT if tipo == 'entrada' else 'FFFFFF')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BORDA
        ws.column_dimensions[get_column_letter(i)].width = larg
    ws.row_dimensions[linha].height = 30
    ws.freeze_panes = ws.cell(row=linha + 1, column=1)


def corpo(ws, colunas, primeira, ultima, formatos=None):
    """Pinta a área de lançamento: creme onde se digita, azul-claro onde calcula."""
    formatos = formatos or {}
    for r in range(primeira, ultima + 1):
        for i, (_, _, tipo) in enumerate(colunas, start=1):
            c = ws.cell(row=r, column=i)
            c.border = BORDA
            c.fill = PatternFill('solid', fgColor=ENTRADA if tipo == 'entrada' else CALC)
            c.font = Font(name='Calibri', size=10,
                          color=INK if tipo == 'entrada' else NAVY2,
                          italic=(tipo != 'entrada'))
            if i in formatos:
                c.number_format = formatos[i]


def band(ws, texto, linha, ate, cor=NAVY):
    ws.merge_cells(start_row=linha, start_column=2, end_row=linha, end_column=ate)
    c = ws.cell(row=linha, column=2, value=texto)
    c.fill = PatternFill('solid', fgColor=cor)
    c.font = Font(name='Calibri', size=9, bold=True, color=GOLDSOFT)
    c.alignment = Alignment(horizontal='left', vertical='center', indent=1)
    ws.row_dimensions[linha].height = 20


# ══════════════════════════════════════════════════════════════════════════════
wb = Workbook()

# ── 1 · INSTRUÇÕES ────────────────────────────────────────────────────────────
ws = wb.active
ws.title = 'Instruções'
ws.sheet_properties.tabColor = GOLD
titulo(ws, 'GESTÃO MENSAL 2026 — COMO USAR', 'Leia uma vez. As definições daqui é que impedem o número de mudar conforme quem preenche.', 6)
for col, larg in zip('ABCDEF', [3, 26, 62, 32, 20, 20]):
    ws.column_dimensions[col].width = larg

L = 4


def txt(linha, col, valor, *, negrito=False, cor=INK, tam=10, wrap=True, italico=False):
    c = ws.cell(row=linha, column=col, value=valor)
    c.font = Font(name='Calibri', size=tam, bold=negrito, color=cor, italic=italico)
    c.alignment = Alignment(vertical='top', wrap_text=wrap)
    return c


band(ws, 'O QUE ESTA PLANILHA RESPONDE', L, 6); L += 1
txt(L, 2, 'A pergunta do Walton', negrito=True, cor=NAVY)
txt(L, 3, 'O que aconteceu na empresa em cada mês, em regime de competência: quanto se vendeu, '
          'quanto de material virou móvel, quanto se fabricou, quanto se instalou e quanto se faturou.')
ws.row_dimensions[L].height = 30; L += 1
txt(L, 2, 'O que ela NÃO responde', negrito=True, cor=RED)
txt(L, 3, 'Se o dinheiro entrou na conta. Isto é competência, não caixa — o mesmo mês dá números '
          'bem diferentes pelos dois critérios. As colunas de recebimento e de parcelas já existem '
          'para o dia em que o Walton pedir a visão de caixa.')
ws.row_dimensions[L].height = 40; L += 2

band(ws, 'AS DEFINIÇÕES — QUANDO CADA COISA CONTA', L, 6); L += 1
for nome, quando, onde in [
    ('R$ vendido', 'No mês da assinatura do contrato. Vale pelo projeto inteiro.', 'aba Projetos'),
    ('Material comprado', 'No mês da compra, pelo valor cheio — parcelamento não divide o valor, '
                          'porque parcela é caixa e não competência.', 'aba Compras'),
    ('R$ fabricado', 'No mês em que o ambiente é liberado no portão da fábrica, conferido '
                     '(POP-01). Ambiente que atravessa o mês conta inteiro no mês em que ficou pronto.', 'aba Ambientes'),
    ('Material aplicado', 'Junto com o R$ fabricado: o material de um ambiente conta no mês em que '
                          'aquele ambiente ficou pronto.', 'calculado'),
    ('R$ instalado', 'No mês em que o ambiente é montado e aceito pelo cliente.', 'aba Ambientes'),
    ('R$ faturado', 'No mês da emissão da nota. Segue o contrato, não a entrega.', 'aba Faturamento'),
    ('Custo fixo', 'Na competência do mês: folha, estrutura e frota.', 'aba Custo Fixo'),
]:
    txt(L, 2, nome, negrito=True, cor=NAVY)
    txt(L, 3, quando)
    txt(L, 4, onde, italico=True, cor=MUTED)
    ws.row_dimensions[L].height = 28
    L += 1
L += 1

band(ws, 'COMO SE PREENCHE', L, 6); L += 1
for n, passo in enumerate([
    'Cadastre o contrato em Projetos: nome, valor e o mês da venda.',
    'Quebre o projeto em ambientes na aba Ambientes, com o valor de cada um. A soma tem de fechar '
    'com o contrato — a coluna Confere avisa quando não fecha.',
    'Conforme a obra anda, marque o Mês fabricado e o Mês instalado de cada ambiente. É só isso '
    'que move o R$ fabricado, o R$ instalado e o material aplicado.',
    'Lance as compras em Compras, uma linha por nota, com o projeto. Compra que não é de projeto '
    'nenhum (fita, cola, parafuso, ferramenta) fica com o projeto em branco e vira material de uso geral.',
    'Lance as notas em Faturamento e o custo fixo do mês em Custo Fixo.',
    'As lâminas de Jan a Dez e a de Ano 2026 se calculam sozinhas. Não digite nada nelas.',
], start=1):
    txt(L, 2, f'Passo {n}', negrito=True, cor=NAVY2)
    txt(L, 3, passo)
    ws.row_dimensions[L].height = 30
    L += 1
L += 1

band(ws, 'O QUE CUIDAR', L, 6, cor=RED); L += 1
for aviso in [
    'Fundo creme é onde se digita. Fundo azul-claro é calculado — não digite por cima, a fórmula se perde.',
    'O nome do projeto tem de ser escrito igual nas quatro abas. Se sair diferente, a coluna Aviso '
    'em Ambientes acende e a soma do projeto não fecha.',
    'O rateio do material por valor é aproximação: cozinha consome mais material por real que um home. '
    'No total do mês se compensa. Quando você souber o material real de um ambiente, informe na coluna '
    'Material apurado — o rateio passa a se aplicar só ao que sobrou.',
    'Ambiente grande demais para caber num mês se quebra em dois ambientes. Não lance percentual: '
    'seria trazer de volta a estimativa que o lançamento por ambiente veio resolver.',
    'Quem informa que o ambiente ficou pronto e que foi aceito precisa ter nome. Sem dono, a planilha '
    'desatualiza em três semanas e vira ficção.',
]:
    txt(L, 3, '•  ' + aviso)
    ws.row_dimensions[L].height = 32
    L += 1

# ── 2 · PROJETOS ──────────────────────────────────────────────────────────────
ws = wb.create_sheet('Projetos')
COLS_P = [('Projeto', 26, 'entrada'), ('Valor do contrato', 18, 'entrada'),
          ('Mês da venda', 13, 'entrada'), ('Origem / vendedor', 18, 'entrada'),
          ('Status', 15, 'entrada'), ('Σ ambientes', 16, 'calc'),
          ('Confere?', 22, 'calc'), ('nº mês', 8, 'calc')]
titulo(ws, 'PROJETOS — um contrato por linha', 'Só o que identifica e o valor. O detalhe vai para a aba Ambientes.', len(COLS_P))
cabecalho(ws, COLS_P)
corpo(ws, COLS_P, 4, 3 + LINHAS_PROJETOS, {2: DINHEIRO, 6: DINHEIRO})
for r in range(4, 4 + LINHAS_PROJETOS):
    ws.cell(row=r, column=6, value=f'=IF($A{r}="","",SUMIF(Ambientes!$A:$A,$A{r},Ambientes!$C:$C))')
    ws.cell(row=r, column=7, value=(
        f'=IF($A{r}="","",IF($B{r}="","informe o valor do contrato",'
        f'IF(ABS($B{r}-$F{r})<1,"OK","difere em "&TEXT($B{r}-$F{r},"R$ #,##0"))))'))
    ws.cell(row=r, column=8, value=f'=IF($C{r}="","",IFERROR(MATCH($C{r},Listas!$A$4:$A$15,0),""))')
menu(ws, MESES, f'C4:C{3 + LINHAS_PROJETOS}')
menu(ws, VENDEDORES, f'D4:D{3 + LINHAS_PROJETOS}')
menu(ws, STATUS, f'E4:E{3 + LINHAS_PROJETOS}')
ws.conditional_formatting.add(
    f'G4:G{3 + LINHAS_PROJETOS}',
    FormulaRule(formula=[f'AND($G4<>"",$G4<>"OK")'],
                fill=PatternFill('solid', fgColor='FBE3E2'), font=Font(color=RED, bold=True)))
ws.column_dimensions['H'].hidden = True

# ── 3 · AMBIENTES ─────────────────────────────────────────────────────────────
ws = wb.create_sheet('Ambientes')
ws.sheet_properties.tabColor = RED
COLS_A = [('Projeto', 24, 'entrada'), ('Ambiente', 26, 'entrada'), ('Valor do ambiente', 17, 'entrada'),
          ('Mês fabricado', 14, 'entrada'), ('Mês instalado', 14, 'entrada'),
          ('Situação', 14, 'calc'), ('Material apurado\n(opcional)', 16, 'entrada'),
          ('Material do ambiente', 17, 'calc'), ('Aviso', 24, 'calc'), ('nº mês', 8, 'calc')]
titulo(ws, 'AMBIENTES — o coração da planilha',
       'Um ambiente por linha. É o Mês fabricado e o Mês instalado que movem quase todos os números.', len(COLS_A))
cabecalho(ws, COLS_A)
corpo(ws, COLS_A, 4, 3 + LINHAS_AMBIENTES, {3: DINHEIRO, 7: DINHEIRO, 8: DINHEIRO})
ult = 3 + LINHAS_AMBIENTES
for r in range(4, ult + 1):
    ws.cell(row=r, column=6, value=(
        f'=IF($A{r}="","",IF($E{r}<>"","Instalado",IF($D{r}<>"","Fabricado","A produzir")))'))
    # material do ambiente: apurado, se houver; senão o rateio do que sobrou das compras do projeto
    base = (f'(SUMIF($A$4:$A${ult},$A{r},$C$4:$C${ult})'
            f'-SUMIFS($C$4:$C${ult},$A$4:$A${ult},$A{r},$G$4:$G${ult},">0"))')
    bolo = (f'(SUMIF(Compras!$A:$A,$A{r},Compras!$D:$D)'
            f'-SUMIFS($G$4:$G${ult},$A$4:$A${ult},$A{r},$G$4:$G${ult},">0"))')
    ws.cell(row=r, column=8, value=(
        f'=IF($A{r}="","",IF($G{r}>0,$G{r},IF({base}<=0,0,MAX(0,{bolo})*$C{r}/{base})))'))
    ws.cell(row=r, column=9, value=(
        f'=IF($A{r}="","",IF(COUNTIF(Projetos!$A:$A,$A{r})=0,"projeto não cadastrado",""))'))
    ws.cell(row=r, column=10, value=f'=IF($D{r}="","",IFERROR(MATCH($D{r},Listas!$A$4:$A$15,0),""))')
menu(ws, MESES, f'D4:D{ult}')
menu(ws, MESES, f'E4:E{ult}')
ws.conditional_formatting.add(
    f'I4:I{ult}',
    FormulaRule(formula=['$I4<>""'], fill=PatternFill('solid', fgColor='FBE3E2'),
                font=Font(color=RED, bold=True)))
ws.column_dimensions['J'].hidden = True

# ── 4 · FATURAMENTO ───────────────────────────────────────────────────────────
ws = wb.create_sheet('Faturamento')
COLS_F = [('Projeto', 24, 'entrada'), ('Nº da nota', 14, 'entrada'), ('Valor', 16, 'entrada'),
          ('Mês da emissão', 15, 'entrada'), ('Mês do recebimento\n(para caixa)', 17, 'entrada'),
          ('Observação', 34, 'entrada')]
titulo(ws, 'FATURAMENTO — uma nota por linha',
       'O mês da emissão é o que conta na competência. O do recebimento fica guardado para a visão de caixa.', len(COLS_F))
cabecalho(ws, COLS_F)
corpo(ws, COLS_F, 4, 3 + LINHAS_FATURAMENTO, {3: DINHEIRO})
menu(ws, MESES, f'D4:D{3 + LINHAS_FATURAMENTO}')
menu(ws, MESES, f'E4:E{3 + LINHAS_FATURAMENTO}')

# ── 5 · COMPRAS ───────────────────────────────────────────────────────────────
ws = wb.create_sheet('Compras')
COLS_C = [('Projeto\n(em branco = uso geral)', 24, 'entrada'), ('Fornecedor', 22, 'entrada'),
          ('Categoria', 18, 'entrada'), ('Valor', 15, 'entrada'), ('Mês da compra', 14, 'entrada'),
          ('Parcelas\n(para caixa)', 12, 'entrada'), ('Observação', 30, 'entrada'), ('nº mês', 8, 'calc')]
titulo(ws, 'COMPRAS — uma nota por linha',
       'Valor cheio, mesmo parcelado. Sem projeto = material de uso geral (fita, cola, parafuso, ferramenta).', len(COLS_C))
cabecalho(ws, COLS_C)
corpo(ws, COLS_C, 4, 3 + LINHAS_COMPRAS, {4: DINHEIRO})
for r in range(4, 4 + LINHAS_COMPRAS):
    ws.cell(row=r, column=8, value=f'=IF($E{r}="","",IFERROR(MATCH($E{r},Listas!$A$4:$A$15,0),""))')
menu(ws, CATEGORIAS, f'C4:C{3 + LINHAS_COMPRAS}')
menu(ws, MESES, f'E4:E{3 + LINHAS_COMPRAS}')
ws.column_dimensions['H'].hidden = True

# ── 6 · CUSTO FIXO ────────────────────────────────────────────────────────────
ws = wb.create_sheet('Custo Fixo')
COLS_CF = [('Mês', 12, 'calc'), ('Folha', 16, 'entrada'), ('Estrutura', 16, 'entrada'),
           ('Frota', 16, 'entrada'), ('Total', 16, 'calc'), ('Observação', 40, 'entrada')]
titulo(ws, 'CUSTO FIXO — um mês por linha',
       'As mesmas três rubricas que já apresentamos ao Walton em agosto, para os números baterem entre os documentos.', len(COLS_CF))
cabecalho(ws, COLS_CF)
corpo(ws, COLS_CF, 4, 15, {2: DINHEIRO, 3: DINHEIRO, 4: DINHEIRO, 5: DINHEIRO})
for i, m in enumerate(MESES):
    r = 4 + i
    ws.cell(row=r, column=1, value=m).font = Font(name='Calibri', size=10, bold=True, color=NAVY2)
    ws.cell(row=r, column=1).alignment = Alignment(horizontal='center')
    ws.cell(row=r, column=5, value=f'=IF(COUNT($B{r}:$D{r})=0,"",SUM($B{r}:$D{r}))')

# ── 7 · AS DOZE LÂMINAS ───────────────────────────────────────────────────────
TILES = [
    ('CUSTO FIXO', "=SUMIF('Custo Fixo'!$A:$A,\"{m}\",'Custo Fixo'!$E:$E)", 'folha · estrutura · frota', NAVY),
    ('R$ VENDIDO', '=SUMIF(Projetos!$C:$C,"{m}",Projetos!$B:$B)', 'contratos assinados', NAVY),
    ('MATERIAL APLICADO', '=SUMIF(Ambientes!$D:$D,"{m}",Ambientes!$H:$H)', 'dos ambientes fabricados', ROXO),
    ('R$ FABRICADO', '=SUMIF(Ambientes!$D:$D,"{m}",Ambientes!$C:$C)', 'ambientes liberados', NAVY),
    ('R$ INSTALADO', '=SUMIF(Ambientes!$E:$E,"{m}",Ambientes!$C:$C)', 'ambientes aceitos', NAVY),
    ('R$ FATURADO', '=SUMIF(Faturamento!$D:$D,"{m}",Faturamento!$C:$C)', 'notas emitidas', NAVY),
]

for idx, m in enumerate(MESES, start=1):
    ws = wb.create_sheet(m)
    ws.sheet_properties.tabColor = NAVY2
    for col, larg in zip('ABCDEFG', [3, 24, 24, 24, 24, 24, 24]):
        ws.column_dimensions[col].width = larg
    titulo(ws, f'{MES_EXTENSO[m].upper()} / {ANO}',
           'Tudo calculado. Não digite nada nesta aba — os números vêm das abas de lançamento.', 7)

    band(ws, 'OS NÚMEROS DO MÊS  ·  competência', 4, 7)
    for i, (rot, formula, obs, cor) in enumerate(TILES):
        c = 2 + i
        a = ws.cell(row=5, column=c, value=rot)
        a.fill = PatternFill('solid', fgColor=cor)
        a.font = Font(name='Calibri', size=8, bold=True, color=GOLDSOFT)
        a.alignment = Alignment(horizontal='center', vertical='center')
        v = ws.cell(row=6, column=c, value=formula.format(m=m))
        v.fill = PatternFill('solid', fgColor=cor)
        v.font = Font(name='Calibri', size=18, bold=True, color='FFFFFF')
        v.alignment = Alignment(horizontal='center', vertical='center')
        v.number_format = DINHEIRO
        o = ws.cell(row=7, column=c, value=obs)
        o.fill = PatternFill('solid', fgColor=cor)
        o.font = Font(name='Calibri', size=8, color='C6CFDA')
        o.alignment = Alignment(horizontal='center', vertical='top')
    ws.row_dimensions[5].height = 16
    ws.row_dimensions[6].height = 34
    ws.row_dimensions[7].height = 15

    band(ws, 'LEITURAS DO MÊS', 9, 7, cor=NAVY2)
    ate = f'"<={idx}"'
    geral = (f'SUMIF(Compras!$E:$E,"{m}",Compras!$D:$D)'
             f'-SUMIFS(Compras!$D:$D,Compras!$E:$E,"{m}",Compras!$A:$A,"<>")')
    leituras = [
        ('Material comprado', 'compras do mês',
         f'=SUMIF(Compras!$E:$E,"{m}",Compras!$D:$D)', DINHEIRO,
         f'=IF({geral}=0,"tudo vinculado a projeto","dos quais "&TEXT({geral},"R$ #,##0")&" de uso geral")',
         ROXO),
        # o estoque só olha compra vinculada a projeto: material de uso geral (fita, cola,
        # parafuso) nunca é aplicado a ambiente nenhum e ficaria acumulando para sempre,
        # fingindo estoque que não existe
        ('Material em estoque', 'comprado de projeto − aplicado, acumulado',
         f'=SUMIFS(Compras!$D:$D,Compras!$H:$H,{ate},Compras!$A:$A,"<>")'
         f'-SUMIF(Ambientes!$J:$J,{ate},Ambientes!$H:$H)',
         DINHEIRO, '"explica caixa curto com resultado bom"', BLUE),
        ('Material sobre o fabricado', 'aplicado ÷ fabricado',
         '=IF($E$6=0,"",$D$6/$E$6)', PCT, '"o termômetro da margem bruta"', OK),
        ('Parado na fábrica', 'fabricado − instalado', '=$E$6-$F$6', DINHEIRO,
         '"móvel pronto esperando obra"', RED),
        ('A faturar', 'instalado − faturado', '=$F$6-$G$6', DINHEIRO,
         '"entregue e ainda sem nota"', RED),
        ('Carteira a produzir', 'vendido acum. − fabricado acum.',
         f'=SUMIF(Projetos!$H:$H,{ate},Projetos!$B:$B)-SUMIF(Ambientes!$J:$J,{ate},Ambientes!$C:$C)',
         DINHEIRO, '"trabalho que a fábrica já tem contratado"', BLUE),
    ]
    for i, (nome, conta, formula, fmt, pergunta, cor) in enumerate(leituras):
        c = 2 + i
        t = ws.cell(row=10, column=c, value=nome)
        t.font = Font(name='Calibri', size=9, bold=True, color=NAVY)
        t.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        t.border = BORDA
        k = ws.cell(row=11, column=c, value=conta)
        k.font = Font(name='Calibri', size=8, italic=True, color=MUTED)
        k.alignment = Alignment(horizontal='center', wrap_text=True)
        k.border = BORDA
        v = ws.cell(row=12, column=c, value=formula)
        v.font = Font(name='Calibri', size=13, bold=True, color=cor)
        v.alignment = Alignment(horizontal='center', vertical='center')
        v.number_format = fmt
        v.border = BORDA
        v.fill = PatternFill('solid', fgColor=CALC)
        q = ws.cell(row=13, column=c, value='=' + pergunta)
        q.font = Font(name='Calibri', size=8, color=MUTED)
        q.alignment = Alignment(horizontal='center', vertical='top', wrap_text=True)
        q.border = BORDA
    ws.row_dimensions[10].height = 26
    ws.row_dimensions[12].height = 24
    ws.row_dimensions[13].height = 26

    band(ws, 'ONDE ESTE MÊS ACONTECEU  ·  por projeto', 15, 7, cor=NAVY2)
    for i, nome in enumerate(['Projeto', 'Fabricado no mês', 'Instalado no mês',
                              'Material aplicado', 'Faturado no mês', 'Vendido no mês']):
        c = ws.cell(row=16, column=2 + i, value=nome)
        c.fill = PatternFill('solid', fgColor=BLUE)
        c.font = Font(name='Calibri', size=9, bold=True, color='FFFFFF')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = BORDA
    ws.row_dimensions[16].height = 26
    for i in range(LINHAS_PROJETOS):
        r, pr = 17 + i, 4 + i
        ws.cell(row=r, column=2, value=f'=IF(Projetos!$A{pr}="","",Projetos!$A{pr})')
        ws.cell(row=r, column=3, value=f'=IF($B{r}="","",SUMIFS(Ambientes!$C:$C,Ambientes!$A:$A,$B{r},Ambientes!$D:$D,"{m}"))')
        ws.cell(row=r, column=4, value=f'=IF($B{r}="","",SUMIFS(Ambientes!$C:$C,Ambientes!$A:$A,$B{r},Ambientes!$E:$E,"{m}"))')
        ws.cell(row=r, column=5, value=f'=IF($B{r}="","",SUMIFS(Ambientes!$H:$H,Ambientes!$A:$A,$B{r},Ambientes!$D:$D,"{m}"))')
        ws.cell(row=r, column=6, value=f'=IF($B{r}="","",SUMIFS(Faturamento!$C:$C,Faturamento!$A:$A,$B{r},Faturamento!$D:$D,"{m}"))')
        ws.cell(row=r, column=7, value=f'=IF($B{r}="","",SUMIFS(Projetos!$B:$B,Projetos!$A:$A,$B{r},Projetos!$C:$C,"{m}"))')
        for c in range(2, 8):
            cel = ws.cell(row=r, column=c)
            cel.border = BORDA
            cel.font = Font(name='Calibri', size=10, color=NAVY2)
            cel.fill = PatternFill('solid', fgColor=ZEBRA if i % 2 else 'FFFFFF')
            if c > 2:
                cel.number_format = DINHEIRO
    ws.freeze_panes = 'B17'

# ── 8 · ANO ───────────────────────────────────────────────────────────────────
ws = wb.create_sheet(f'Ano {ANO}')
ws.sheet_properties.tabColor = GOLD
ws.column_dimensions['A'].width = 3
ws.column_dimensions['B'].width = 30
for i in range(12):
    ws.column_dimensions[get_column_letter(3 + i)].width = 13
ws.column_dimensions[get_column_letter(15)].width = 15
titulo(ws, f'ANO {ANO} — os doze meses lado a lado',
       'É aqui que se lê tendência. Tudo calculado a partir das lâminas mensais.', 15)

for i, m in enumerate(MESES):
    c = ws.cell(row=4, column=3 + i, value=m)
    c.fill = PatternFill('solid', fgColor=NAVY)
    c.font = Font(name='Calibri', size=10, bold=True, color=GOLDSOFT)
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = BORDA
for col, rot in [(2, 'R$ · competência'), (15, 'Ano')]:
    c = ws.cell(row=4, column=col, value=rot)
    c.fill = PatternFill('solid', fgColor=NAVY)
    c.font = Font(name='Calibri', size=10, bold=True, color=GOLDSOFT)
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = BORDA
ws.row_dimensions[4].height = 24

LINHAS_ANO = [
    ('Custo fixo', 'B6', DINHEIRO, False),
    ('R$ vendido', 'C6', DINHEIRO, False),
    ('Material aplicado', 'D6', DINHEIRO, False),
    ('R$ fabricado', 'E6', DINHEIRO, False),
    ('R$ instalado', 'F6', DINHEIRO, False),
    ('R$ faturado', 'G6', DINHEIRO, False),
    (None, None, None, None),
    ('Material comprado', 'B12', DINHEIRO, True),
    ('Material em estoque', 'C12', DINHEIRO, True),
    ('Material sobre o fabricado', 'D12', PCT, True),
    ('Parado na fábrica', 'E12', DINHEIRO, True),
    ('A faturar', 'F12', DINHEIRO, True),
    ('Carteira a produzir', 'G12', DINHEIRO, True),
]
r = 5
linha_de = {}
for rot, celula, fmt, leitura in LINHAS_ANO:
    if rot is None:
        ws.row_dimensions[r].height = 8
        r += 1
        continue
    linha_de[rot] = r
    c = ws.cell(row=r, column=2, value=rot)
    c.font = Font(name='Calibri', size=10, bold=not leitura, color=BLUE if leitura else NAVY2,
                  italic=leitura)
    c.alignment = Alignment(vertical='center', indent=1)
    c.border = BORDA
    c.fill = PatternFill('solid', fgColor='EDF3F9' if leitura else GOLDBG)
    for i, m in enumerate(MESES):
        v = ws.cell(row=r, column=3 + i, value=f"='{m}'!${celula[0]}${celula[1:]}")
        v.number_format = fmt
        v.font = Font(name='Calibri', size=10, color=BLUE if leitura else INK, italic=leitura)
        v.alignment = Alignment(horizontal='right')
        v.border = BORDA
        v.fill = PatternFill('solid', fgColor='F7FAFD' if leitura else 'FFFFFF')
    total = ws.cell(row=r, column=15)
    if fmt == PCT:
        # no ano é o total aplicado sobre o total fabricado — não a média dos meses,
        # que daria peso igual a um mês de R$ 20 mil e a outro de R$ 200 mil
        ap, fab = linha_de['Material aplicado'], linha_de['R$ fabricado']
        total.value = (f'=IF($O${fab}=0,"",$O${ap}/$O${fab})')
    elif leitura and rot in ('Material em estoque', 'Carteira a produzir'):
        total.value = f'=$N{r}'          # posição acumulada = a do último mês, não a soma
    elif leitura:
        total.value = f'=SUM($C{r}:$N{r})'
    else:
        total.value = f'=SUM($C{r}:$N{r})'
    total.number_format = fmt
    total.font = Font(name='Calibri', size=10, bold=True, color=NAVY)
    total.alignment = Alignment(horizontal='right')
    total.border = BORDA
    total.fill = PatternFill('solid', fgColor=GOLDBG)
    ws.row_dimensions[r].height = 20
    r += 1

nota = ws.cell(row=r + 1, column=2, value=(
    'Material em estoque e carteira a produzir são posições acumuladas — mostram onde a empresa está '
    'no fim de cada mês, não quanto passou no mês. Por isso a coluna Ano repete o valor de dezembro '
    'em vez de somar. "Material sobre o fabricado" no ano é o total aplicado dividido pelo total '
    'fabricado, e não a média dos doze meses, que daria o mesmo peso a um mês de R$ 20 mil e a outro '
    'de R$ 200 mil.'))
nota.font = Font(name='Calibri', size=9, italic=True, color=MUTED)
nota.alignment = Alignment(vertical='top', wrap_text=True)
ws.merge_cells(start_row=r + 1, start_column=2, end_row=r + 2, end_column=15)

# ── 9 · LISTAS ────────────────────────────────────────────────────────────────
ws = wb.create_sheet('Listas')
ws.sheet_properties.tabColor = MUTED
titulo(ws, 'LISTAS',
       'A coluna Meses alimenta o cálculo do nº do mês (não apague nem reordene). '
       'As demais documentam os menus suspensos.', 4)
for col, larg in zip('ABCD', [14, 20, 22, 24]):
    ws.column_dimensions[col].width = larg

for i, (nome, itens) in enumerate([('Meses', MESES), ('Status', STATUS),
                                   ('Origem', VENDEDORES), ('Categorias', CATEGORIAS)]):
    c = ws.cell(row=3, column=1 + i, value=nome)
    c.fill = PatternFill('solid', fgColor=NAVY)
    c.font = Font(name='Calibri', size=9, bold=True, color=GOLDSOFT)
    c.alignment = Alignment(horizontal='center', vertical='center')
    c.border = BORDA
    for j, it in enumerate(itens):
        v = ws.cell(row=4 + j, column=1 + i, value=it)
        v.font = Font(name='Calibri', size=10, color=INK)
        v.border = BORDA
        v.fill = PatternFill('solid', fgColor=CINZA if i == 0 else 'FFFFFF')
ws.row_dimensions[3].height = 20

wb.save(ARQUIVO)
print(f'planilha gerada: {ARQUIVO}')
print(f'  abas: {len(wb.sheetnames)} — {", ".join(wb.sheetnames)}')
