#!/usr/bin/env python3
"""Reconstrói a aba Base · Comissões dentro do arquivo que já está em uso.

    python3 reconstruir-base.py ENTRADA.xlsx SAIDA.xlsx

Para que serve: a Custo por Projeto está viva no Drive e recebe lançamento
todo dia. Quando a Base precisa mudar — coluna nova, projeto novo, projeto
apagado — regerar o arquivo inteiro e migrar os dados de volta apagaria o que
foi digitado desde a última vez. Este script faz o contrário: abre o arquivo
real e reescreve SÓ a aba Base, que é cem por cento fórmula e não guarda nada.

É a mesma rotina do botão "Reconstruir a Base · Comissões" do Apps Script,
numa versão que roda aqui.

O que ele faz:
  · lê os projetos registrados no Painel Geral;
  · limpa do Painel Geral o registro de projeto cuja aba não existe mais;
  · recria a Base com um bloco de doze linhas por projeto que de fato existe.

Nada mais é tocado.
"""
import sys

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

NAVY2, GOLD, LINE = '16314F', 'C2A05A', 'D6DBE2'
MOEDA, PCT1, DATA = 'R$ #,##0.00', '0.0%', 'DD/MM/YYYY'
F = 'Calibri'

ABA_BASE = 'Base · Comissões'
ABA_PAINEL = 'Painel Geral'
PG0, PGF = 11, 50          # faixa de projetos no Painel Geral
AMB0, N_AMB = 49, 12       # bloco de ambientes na ficha
LIN0 = 2                   # primeira linha de dados da Base
R_ID2, R_ID4 = 7, 9        # data de entrada · coordenador
R_COORD, R_LIQ, R_KPI = 65, 45, 13
COLS_AMB = list('ABCDEFGHIJKLMNO')

CAB = [
    ('Projeto', 22, None), ('Cliente', 24, None),
    ('Ambiente', 22, None), ('% do total', 9, PCT1), ('Valor do ambiente', 15, MOEDA),
    ('Produção — quem', 16, None), ('% produção', 9, PCT1),
    ('Comissão produção', 15, MOEDA), ('Fabricado em', 12, DATA),
    ('Montagem — quem', 16, None), ('% montagem', 9, PCT1),
    ('Comissão montagem', 15, MOEDA), ('Montado em', 12, DATA),
    ('Projetista — quem', 16, None), ('% projeto', 9, PCT1),
    ('Comissão projetista', 15, MOEDA), ('Total do ambiente', 15, MOEDA),
    ('Data de entrada', 12, DATA),
    ('Coordenador', 16, None), ('Comissão coordenação', 15, MOEDA),
    ('Receita líquida do projeto', 16, MOEDA), ('Valor de venda', 15, MOEDA),
]


def ref(nome):
    return "'" + str(nome).replace("'", "''") + "'"


def main(entrada, saida):
    wb = openpyxl.load_workbook(entrada)
    pg = wb[ABA_PAINEL]

    # ── quem está registrado e ainda existe ──
    projetos, orfaos = [], []
    for r in range(PG0, PGF + 1):
        nome = pg.cell(r, 2).value
        if not nome:
            continue
        if nome in wb.sheetnames:
            projetos.append((r, nome))
        else:
            orfaos.append((r, nome))
            pg.cell(r, 2).value = None

    # ── Base do zero ──
    if ABA_BASE in wb.sheetnames:
        idx = wb.sheetnames.index(ABA_BASE)
        del wb[ABA_BASE]
    else:
        idx = len(wb.sheetnames)
    bs = wb.create_sheet(ABA_BASE, idx)
    bs.sheet_properties.tabColor = GOLD
    bs.sheet_view.showGridLines = False
    bs.freeze_panes = 'C2'
    for i, (titulo, larg, _) in enumerate(CAB, start=1):
        c = bs.cell(1, i, titulo)
        c.font = Font(name=F, size=8.5, bold=True, color='FFFFFF')
        c.fill = PatternFill('solid', fgColor=NAVY2)
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        c.border = Border(bottom=Side('medium', color=GOLD))
        bs.column_dimensions[get_column_letter(i)].width = larg
    bs.row_dimensions[1].height = 30

    lin = LIN0
    for pg_row, nome in projetos:
        ab = ref(nome)
        for j in range(N_AMB):
            fr = AMB0 + j
            bs.cell(lin, 1, f'=IF($C{lin}="","",\'{ABA_PAINEL}\'!$B${pg_row})')
            bs.cell(lin, 2, f'=IF($C{lin}="","",\'{ABA_PAINEL}\'!$C${pg_row})')
            for k, cl in enumerate(COLS_AMB):
                bs.cell(lin, 3 + k, f'=IF({ab}!${cl}${fr}="","",{ab}!${cl}${fr})')
            bs.cell(lin, 18, f'=IF($C{lin}="","",IF({ab}!$F${R_ID2}="","",{ab}!$F${R_ID2}))')
            if j == 0:      # coordenação e totais do projeto, uma vez por bloco
                bs.cell(lin, 19, f'=IF($C{lin}="","",IF({ab}!$F${R_ID4}="","",{ab}!$F${R_ID4}))')
                bs.cell(lin, 20, f'=IF($C{lin}="","",IF({ab}!$D${R_COORD}="",0,{ab}!$D${R_COORD}))')
                bs.cell(lin, 21, f'=IF($C{lin}="","",IF({ab}!$D${R_LIQ}="",0,{ab}!$D${R_LIQ}))')
                bs.cell(lin, 22, f'=IF($C{lin}="","",IF({ab}!$A${R_KPI}="",0,{ab}!$A${R_KPI}))')
            for k, (_, _, nf) in enumerate(CAB, start=1):
                cel = bs.cell(lin, k)
                cel.font = Font(name=F, size=9)
                if nf:
                    cel.number_format = nf
                cel.border = Border(bottom=Side('thin', color=LINE))
            lin += 1

    wb.save(saida)
    print('═' * 66)
    print('BASE · COMISSÕES reconstruída')
    print('═' * 66)
    print(f'  projetos .... {len(projetos)}')
    print(f'  linhas ...... {lin - LIN0}')
    print(f'  colunas ..... {len(CAB)}')
    if orfaos:
        print('\n  REGISTRO LIMPO — estava no Painel Geral sem aba correspondente:')
        for r, n in orfaos:
            print(f'    linha {r}: {n!r}')
    print(f'\n  abas preservadas: {len(wb.sheetnames)}')
    print(f'\nOK → {saida}')


if __name__ == '__main__':
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
