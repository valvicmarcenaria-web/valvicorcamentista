#!/usr/bin/env python3
"""
Testa a planilha de gestão mensal do Walton — de verdade: preenche dados de exemplo,
manda o LibreOffice recalcular e confere os números contra o que a conta dá na mão.

    python3 testar-gestao-mensal.py

Sem isso, a planilha só teria sido "gerada sem erro" — o que não diz nada sobre as
fórmulas estarem certas. As contas que este teste prova são as que sustentam o
documento inteiro: o rateio do material, o material aplicado por mês de fabricação,
os acumulados e a conferência Σ ambientes × contrato.
"""
import os
import shutil
import subprocess
import sys
import tempfile

from openpyxl import load_workbook

AQUI = os.path.dirname(os.path.abspath(__file__))
ARQUIVO = os.path.join(AQUI, 'Valvic_Gestao_Mensal_2026.xlsx')
falhas, checagens = [], 0


def confere(rotulo, obtido, esperado, tol=0.51):
    global checagens
    checagens += 1
    if esperado is None:
        ok = obtido in (None, '', 0)
    elif isinstance(esperado, str):
        ok = str(obtido).strip() == esperado
    elif isinstance(esperado, bool):
        ok = obtido is esperado
    else:
        ok = obtido is not None and not isinstance(obtido, str) and abs(obtido - esperado) <= tol
    if not ok:
        falhas.append(f'{rotulo}: obtido {obtido!r}, esperado {esperado!r}')
    return ok


# ── 1 · preenche o exemplo ────────────────────────────────────────────────────
wb = load_workbook(ARQUIVO)

p = wb['Projetos']
for i, (nome, valor, mes) in enumerate([('Alfa', 100000, 'Set'), ('Beta', 60000, 'Ago')]):
    p.cell(row=4 + i, column=1, value=nome)
    p.cell(row=4 + i, column=2, value=valor)
    p.cell(row=4 + i, column=3, value=mes)

a = wb['Ambientes']
for i, (proj, amb, val, fab, inst, apurado) in enumerate([
        ('Alfa', 'Cozinha', 40000, 'Set', 'Set', None),
        ('Alfa', 'Roupeiro', 35000, 'Set', 'Out', None),
        ('Alfa', 'Home', 25000, 'Out', None, None),
        ('Beta', 'Cozinha', 60000, 'Ago', 'Ago', 20000)]):   # apurado, não rateado
    r = 4 + i
    a.cell(row=r, column=1, value=proj)
    a.cell(row=r, column=2, value=amb)
    a.cell(row=r, column=3, value=val)
    a.cell(row=r, column=4, value=fab)
    if inst:
        a.cell(row=r, column=5, value=inst)
    if apurado:
        a.cell(row=r, column=7, value=apurado)

c = wb['Compras']
for i, (proj, valor, mes) in enumerate([('Alfa', 30000, 'Set'), ('Alfa', 15000, 'Out'),
                                        ('Beta', 25000, 'Ago'), (None, 5000, 'Set')]):
    r = 4 + i
    if proj:
        c.cell(row=r, column=1, value=proj)
    c.cell(row=r, column=4, value=valor)
    c.cell(row=r, column=5, value=mes)

f = wb['Faturamento']
for i, (proj, valor, mes) in enumerate([('Alfa', 50000, 'Set'), ('Beta', 60000, 'Ago')]):
    f.cell(row=4 + i, column=1, value=proj)
    f.cell(row=4 + i, column=3, value=valor)
    f.cell(row=4 + i, column=4, value=mes)

cf = wb['Custo Fixo']
for linha, (folha, estrut, frota) in {11: (45000, 20000, 5000),    # Ago → 70.000
                                      12: (50000, 22000, 5000)}.items():  # Set → 77.000
    cf.cell(row=linha, column=2, value=folha)
    cf.cell(row=linha, column=3, value=estrut)
    cf.cell(row=linha, column=4, value=frota)

tmp = tempfile.mkdtemp(prefix='teste-gestao-')
entrada = os.path.join(tmp, 'entrada.xlsx')
wb.save(entrada)

# ── 2 · o LibreOffice recalcula ───────────────────────────────────────────────
# Por padrão o LibreOffice NÃO recalcula fórmulas ao abrir um .xlsx — devolveria o
# arquivo com as células vazias e o teste passaria a testar nada. Por isso o perfil é
# criado do zero e a opção OOXMLRecalcMode vai para 0 ("recalcular sempre").
perfil = os.path.join(tmp, 'perfil')
os.makedirs(perfil, exist_ok=True)
subprocess.run(['libreoffice', '--headless', '--terminate_after_init'],
               capture_output=True, timeout=300, env=dict(os.environ, HOME=perfil))
xcu = os.path.join(perfil, '.config', 'libreoffice', '4', 'user', 'registrymodifications.xcu')
if not os.path.exists(xcu):
    print('não consegui criar o perfil do LibreOffice — sem recálculo o teste não vale')
    sys.exit(1)
conf = open(xcu, encoding='utf-8').read()
conf = conf.replace('</oor:items>',
                    '<item oor:path="/org.openoffice.Office.Calc/Formula/Load">'
                    '<prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop>'
                    '</item></oor:items>')
open(xcu, 'w', encoding='utf-8').write(conf)

destino = os.path.join(tmp, 'recalculado')
r = subprocess.run(['libreoffice', '--headless', '--norestore', '--convert-to', 'xlsx',
                    '--outdir', destino, entrada],
                   capture_output=True, text=True, timeout=600,
                   env=dict(os.environ, HOME=perfil))
recalc = os.path.join(destino, 'entrada.xlsx')
if not os.path.exists(recalc):
    print('LibreOffice não converteu o arquivo:\n' + (r.stderr or r.stdout))
    sys.exit(1)
wb = load_workbook(recalc, data_only=True)

# ── 3 · confere ───────────────────────────────────────────────────────────────
print('\nAmbientes — material do ambiente')
a = wb['Ambientes']
confere('  Alfa · Cozinha   45.000 × 40/100', a['H4'].value, 18000)
confere('  Alfa · Roupeiro  45.000 × 35/100', a['H5'].value, 15750)
confere('  Alfa · Home      45.000 × 25/100', a['H6'].value, 11250)
confere('  Beta · Cozinha   apurado, não rateado', a['H7'].value, 20000)
confere('  situação · instalado', a['F4'].value, 'Instalado')
confere('  situação · fabricado', a['F6'].value, 'Fabricado')
confere('  aviso vazio quando o projeto existe', a['I4'].value, None)

print('Projetos — conferência Σ ambientes × contrato')
p = wb['Projetos']
confere('  Alfa · Σ ambientes', p['F4'].value, 100000)
confere('  Alfa · confere', p['G4'].value, 'OK')
confere('  Beta · confere', p['G5'].value, 'OK')

print('Agosto')
g = wb['Ago']
for rot, cel, esp in [('custo fixo', 'B6', 70000), ('vendido', 'C6', 60000),
                      ('material aplicado', 'D6', 20000), ('fabricado', 'E6', 60000),
                      ('instalado', 'F6', 60000), ('faturado', 'G6', 60000)]:
    confere(f'  {rot}', g[cel].value, esp)

print('Setembro')
s = wb['Set']
for rot, cel, esp in [('custo fixo', 'B6', 77000), ('vendido', 'C6', 100000),
                      ('material aplicado  18.000 + 15.750', 'D6', 33750),
                      ('fabricado          40.000 + 35.000', 'E6', 75000),
                      ('instalado', 'F6', 40000), ('faturado', 'G6', 50000),
                      ('material comprado  30.000 + 5.000 de uso geral', 'B12', 35000),
                      ('material em estoque 55.000 − 53.750', 'C12', 1250),
                      ('parado na fábrica  75.000 − 40.000', 'E12', 35000),
                      ('a faturar          40.000 − 50.000', 'F12', -10000),
                      ('carteira a produzir 160.000 − 135.000', 'G12', 25000)]:
    confere(f'  {rot}', s[cel].value, esp)
confere('  material ÷ fabricado', s['D12'].value, 0.45, tol=0.001)
# o separador de milhar do TEXT() sai no idioma de quem abre o arquivo — aqui o
# LibreOffice roda em inglês e escreve 5,000; no Brasil sairá 5.000. Comparamos sem ele.
confere('  nota do uso geral abaixo do comprado',
        str(s['B13'].value).replace('.', '').replace(',', ''),
        'dos quais R$ 5000 de uso geral')

print('Setembro — quebra por projeto')
for rot, cel, esp in [('linha 1 · projeto', 'B17', 'Alfa'),
                      ('linha 1 · fabricado no mês', 'C17', 75000),
                      ('linha 1 · instalado no mês', 'D17', 40000),
                      ('linha 1 · material aplicado', 'E17', 33750),
                      ('linha 1 · faturado no mês', 'F17', 50000),
                      ('linha 2 · projeto', 'B18', 'Beta'),
                      ('linha 2 · fabricado no mês (nada em set)', 'C18', 0)]:
    confere(f'  {rot}', s[cel].value, esp)

print('Outubro — o que escorreu para o mês seguinte')
o = wb['Out']
for rot, cel, esp in [('fabricado (o Home)', 'E6', 25000),
                      ('instalado (o Roupeiro)', 'F6', 35000),
                      ('material aplicado (o Home)', 'D6', 11250),
                      ('carteira a produzir zerada', 'G12', 0)]:
    confere(f'  {rot}', o[cel].value, esp)

print('Ano 2026 — a lâmina de tendência')
ano = wb['Ano 2026']
for rot, cel, esp in [('fabricado · Ago', 'J8', 60000), ('fabricado · Set', 'K8', 75000),
                      ('fabricado · ano', 'O8', 160000), ('vendido · ano', 'O6', 160000),
                      ('material aplicado · ano', 'O7', 65000),
                      ('material comprado · ano  soma dos doze meses', 'O12', 75000),
                      ('material em estoque · ano = dezembro (70.000 − 65.000)', 'O13', 5000),
                      ('carteira a produzir · ano = posição de dezembro', 'O17', 0)]:
    confere(f'  {rot}', ano[cel].value, esp)
confere('  material sobre o fabricado · ano  65.000 ÷ 160.000', ano['O14'].value, 0.40625, tol=0.001)

print('Menus suspensos — listas literais em todas as abas de lançamento')
wv = load_workbook(ARQUIVO)
for aba, quantos in {'Projetos': 3, 'Ambientes': 2, 'Faturamento': 2, 'Compras': 2}.items():
    dvs = list(wv[aba].data_validations.dataValidation)
    confere(f'  {aba} · nº de menus', len(dvs), quantos)
    for dv in dvs:
        confere(f'  {aba} · lista literal', dv.formula1.startswith('"'), True)
        confere(f'  {aba} · ≤ 255 caracteres', len(dv.formula1) <= 255, True)
        confere(f'  {aba} · setinha visível', dv.showDropDown is False, True)
confere('20 abas', len(wv.sheetnames), 20)

shutil.rmtree(tmp, ignore_errors=True)
print('\n' + '─' * 64)
if falhas:
    print(f'{len(falhas)} FALHA(S) em {checagens} checagens:\n')
    for x in falhas:
        print('  x ' + x)
    sys.exit(1)
print(f'{checagens} checagens · 0 falhas — as contas fecham.')
